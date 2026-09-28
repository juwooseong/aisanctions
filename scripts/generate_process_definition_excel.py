# -*- coding: utf-8 -*-
"""업무프로세스 정의서 — Excel 셀 흐름도 + 상세 흐름 생성."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from excel_native_flowchart import (
    fit_flowchart_columns,
    render_flowchart,
    reset_manifest,
    save_manifest,
)
from generate_screen_design_excel import (
    BORDER,
    SCREENS,
    WRAP,
    auto_fit_columns,
    sheet_title,
    write_kv_table,
)
from process_definition_data import (
    DOMAINS,
    DOMAIN_FLOW,
    build_all_profiles,
    profile_summary,
)

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "process" / "excel"
FLOWCHART_DIR = OUT_DIR / "flowcharts"
MASTER_FILE = OUT_DIR / "업무프로세스정의서_TO-BE.xlsx"
MANIFEST_FILE = OUT_DIR / "flowchart_manifest.json"

PROC_HEADER_FILL = PatternFill("solid", fgColor="7C3AED")
PROC_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
PROC_SECTION_FILL = PatternFill("solid", fgColor="F5F3FF")
PROC_SECTION_FONT = Font(bold=True, size=11, color="5B21B6")
PROC_TAB_COLOR = "7C3AED"


def write_proc_section(ws, row: int, title: str, col_span: int = 10) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    cell = ws.cell(row=row, column=1, value=title)
    cell.fill = PROC_SECTION_FILL
    cell.font = PROC_SECTION_FONT
    cell.alignment = Alignment(vertical="center")
    for c in range(1, col_span + 1):
        ws.cell(row=row, column=c).border = BORDER
    return row + 1


def write_proc_table(ws, row: int, headers: list[str], rows: list[tuple]) -> int:
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = PROC_HEADER_FILL
        c.font = PROC_HEADER_FONT
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row += 1
    for data in rows:
        for col, val in enumerate(data, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = BORDER
            c.alignment = WRAP
        row += 1
    return row + 1


def build_screen_process_sheet(ws, profile: dict) -> None:
    sid = profile["screen_id"]
    bpmn = profile["bpmn_data"]
    summary = profile_summary(profile)
    flowchart_range: tuple[int, int] | None = None

    row = 1
    ws.merge_cells("A1:J1")
    t = ws["A1"]
    t.value = f"업무프로세스 정의서 — SCR-{sid} {profile['name']}"
    t.font = Font(bold=True, size=14, color="5B21B6")
    t.alignment = Alignment(horizontal="center", vertical="center")
    row = 3

    row = write_proc_section(ws, row, "1. 기본 정보")
    row = write_kv_table(ws, row, [
        ("화면ID", sid),
        ("화면명", profile["name"]),
        ("도메인", profile["domain"]),
        ("URL", profile["url"]),
        ("패턴", profile["pattern"]),
        ("상태", profile["status"]),
        ("JSP", f"WEB-INF/jsp/{profile['jsp']}"),
        ("API", profile["api"]),
        ("관련 FR", profile["fr"] or "—"),
        ("비고", profile["notes"] or "—"),
    ])
    row += 1

    row = write_proc_section(ws, row, "2. 프로세스 개요")
    row = write_kv_table(ws, row, [
        ("목적", profile["purpose"]),
        ("권한", profile["auth"][:100] if profile["auth"] else "—"),
        ("유효성", profile["validation"][:100] if profile["validation"] else "—"),
        ("흐름 노드", str(summary.get("흐름노드", 0))),
        ("상세 단계", str(summary.get("상세단계", 0))),
    ])
    row += 1

    row = write_proc_section(ws, row, "3. 선행·종료 조건")
    cond_rows = [(f"선행 {i+1}", p) for i, p in enumerate(profile["pre"])]
    cond_rows += [(f"종료 {i+1}", p) for i, p in enumerate(profile["post"])]
    row = write_proc_table(ws, row, ["구분", "조건"], cond_rows)
    row += 1

    # §4 플로우차트 (Excel COM 후처리로 네이티브 도형·연결선 생성)
    row = write_proc_section(ws, row, "4. 업무 흐름도 (플로우차트)", 16)
    chart_title = f"SCR-{sid} {profile['name']}"
    next_row, flowchart_range = render_flowchart(
        ws, row, bpmn,
        steps_fallback=profile.get("steps"),
        title=chart_title,
        screen_id=sid,
        cache_dir=FLOWCHART_DIR,
    )
    ws._flowchart_range = flowchart_range  # type: ignore[attr-defined]
    fit_flowchart_columns(ws, flowchart_range[0], flowchart_range[1])
    row = next_row + 1

    sec = 5
    steps = profile.get("steps") or []
    if steps:
        row = write_proc_section(ws, row, f"{sec}. 상세 처리 단계")
        row = write_proc_table(
            ws, row,
            ["순번", "단계", "담당", "트리거", "처리내용", "API/함수", "입력", "출력", "분기", "예외"],
            [tuple(s) for s in steps],
        )
        row += 1
        sec += 1

    events = profile.get("events") or []
    if events:
        row = write_proc_section(ws, row, f"{sec}. 이벤트·함수 매핑")
        row = write_proc_table(ws, row, ["이벤트/함수", "설명"], [(e[0], e[1]) for e in events])
        row += 1
        sec += 1

    apis = profile.get("apis") or []
    if apis:
        row = write_proc_section(ws, row, f"{sec}. API 연동")
        row = write_proc_table(ws, row, ["API", "설명"], [(a[0], a[1]) for a in apis])
        row += 1
        sec += 1

    states = profile.get("states") or []
    if states:
        row = write_proc_section(ws, row, f"{sec}. 화면 상태·코드")
        row = write_proc_table(ws, row, ["코드/조건", "설명"], states)
        row += 1
        sec += 1

    ui = profile.get("ui") or []
    if ui:
        row = write_proc_section(ws, row, f"{sec}. UI 구성")
        for item in ui:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            ws.cell(row=row, column=1, value=f"• {item}").alignment = WRAP
            row += 1
        row += 1
        sec += 1

    if profile.get("tobe_changes"):
        row = write_proc_section(ws, row, f"{sec}. TO-BE 변경")
        for ch in profile["tobe_changes"]:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            ws.cell(row=row, column=1, value=f"• {ch}").alignment = WRAP
            row += 1
        row += 1
        sec += 1

    related = profile.get("related") or []
    if related:
        row = write_proc_section(ws, row, f"{sec}. 관련 화면 연계")
        row = write_proc_table(ws, row, ["관계", "화면"], related)
        row += 1
        sec += 1

    tables = profile.get("tables") or []
    if tables:
        row = write_proc_section(ws, row, f"{sec}. 관련 테이블")
        tbl_rows = [(t[0], t[1], t[2] if len(t) > 2 else "") for t in tables[:12]]
        write_proc_table(ws, row, ["논리명", "물리명", "용도"], tbl_rows)

    color = PROC_TAB_COLOR if profile["status"] in ("신규", "변경", "v2") else "94A3B8"
    ws.sheet_properties.tabColor = color


def build_index_sheet(ws, profiles: dict) -> None:
    ws.title = "목차"
    headers = [
        "No", "화면ID", "화면명", "도메인", "패턴", "상태",
        "흐름노드", "단계수", "API수", "시트명",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = PROC_HEADER_FILL
        c.font = PROC_HEADER_FONT
        c.border = BORDER

    widths = [5, 12, 22, 14, 8, 8, 8, 8, 6, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for i, s in enumerate(SCREENS, 1):
        sid, name, domain, *_rest = s
        p = profiles[sid]
        sm = profile_summary(p)
        stitle = sheet_title(sid, name)
        row = i + 1
        ws.cell(row=row, column=1, value=i).border = BORDER
        ws.cell(row=row, column=2, value=sid).border = BORDER
        ws.cell(row=row, column=3, value=name).border = BORDER
        ws.cell(row=row, column=4, value=domain).border = BORDER
        ws.cell(row=row, column=5, value=s[4]).border = BORDER
        ws.cell(row=row, column=6, value=s[5]).border = BORDER
        ws.cell(row=row, column=7, value=sm["흐름노드"]).border = BORDER
        ws.cell(row=row, column=8, value=sm["상세단계"]).border = BORDER
        ws.cell(row=row, column=9, value=sm["API"]).border = BORDER
        c = ws.cell(row=row, column=10, value=stitle)
        c.border = BORDER
        c.hyperlink = f"#'{stitle}'!A1"
        c.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = PROC_TAB_COLOR


def build_domain_sheet(ws, domain_id: str, domain_name: str, screen_ids: list[str], profiles: dict) -> None:
    row = write_proc_section(ws, 1, f"{domain_id} — {domain_name} 도메인 개요", 8)
    row = write_kv_table(ws, row, [
        ("도메인", domain_name),
        ("대표 흐름", DOMAIN_FLOW.get(domain_id, "화면별 독립 프로세스")),
        ("화면 수", str(len(screen_ids))),
    ])
    row += 1

    row = write_proc_section(ws, row, "화면 목록", 8)
    dom_rows = []
    for sid in screen_ids:
        if sid not in profiles:
            continue
        p = profiles[sid]
        sm = profile_summary(p)
        dom_rows.append((
            sid, p["name"], p["pattern"], p["status"],
            p["purpose"][:50], sm["흐름노드"], sm["상세단계"], sm["API"],
            sheet_title(sid, p["name"]),
        ))
    write_proc_table(
        ws, row,
        ["화면ID", "화면명", "패턴", "상태", "목적", "흐름노드", "단계", "API", "시트"],
        dom_rows,
    )
    ws.sheet_properties.tabColor = "8B5CF6"


def build_standard_sheets(wb) -> None:
    ws = wb.create_sheet("00_표준")
    row = write_proc_section(ws, 1, "업무프로세스 정의서 표준", 4)
    write_proc_table(ws, row, ["항목", "내용"], [
        ("문서 목적", "화면별 업무 처리 흐름을 플로우차트(도형+화살표)로 시각화"),
        ("흐름도 방식", "Excel 네이티브 도형 — 타원·둥근사각형·마름모 + Connector 화살표"),
        ("노드 도형", "● 시작(타원)  ■ 종료(타원)  ◆ 분기(마름모)  □ 업무(둥근사각)"),
        ("연결선", "직선 화살표 + 분기 라벨 / ↺ 되돌아감(주황)"),
        ("편집", "Excel에서 개별 도형·연결선·라벨 선택 및 수정 가능"),
        ("갱신", "python scripts/generate_process_definition_excel.py"),
    ])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 56
    ws.sheet_properties.tabColor = PROC_TAB_COLOR

    ws2 = wb.create_sheet("00_흐름도가이드")
    row = write_proc_section(ws2, 1, "플로우차트 도형·화살표", 4)
    write_proc_table(ws2, row, ["도형", "형태", "색상", "설명"], [
        ("시작", "타원(Ellipse)", "#22C55E", "프로세스 진입"),
        ("업무", "둥근 사각형", "#EDE9FE", "API·UI 처리 단계"),
        ("분기", "마름모(Diamond)", "#FEF08A", "조건 분기 — 가로 펼침"),
        ("종료", "타원(Ellipse)", "#64748B", "완료·팝업 닫기"),
        ("화살표", "직선+삼각형 머리", "#7C3AED", "순차·분기 연결"),
        ("되돌아감", "주황 화살표 ↺", "#C2410C", "재조회·루프백"),
    ])
    ws2.sheet_properties.tabColor = PROC_TAB_COLOR

    ws3 = wb.create_sheet("00_역할")
    row = write_proc_section(ws3, 1, "담당자·역할 정의", 4)
    write_proc_table(ws3, row, ["역할", "코드", "대표 화면", "설명"], [
        ("심사자", "S1", "1010, 1011", "본인 배정 건 심사·상신"),
        ("결재자", "S2", "2010, 2011", "교차결재 depth=2 승인/반려"),
        ("QA담당", "QA", "3010, 3011", "자체점검·QA 상신"),
        ("관리자", "mgpeYn=Y", "7050, 8070", "권한·시스템 설정"),
        ("시스템", "—", "전 화면", "자동 처리·API·배치"),
        ("클라이언트", "JS", "전 화면", "UI 이벤트·DataTables"),
    ])
    ws3.sheet_properties.tabColor = PROC_TAB_COLOR


def _auto_fit_sheet(ws) -> None:
    flow_range = getattr(ws, "_flowchart_range", None)
    if flow_range:
        fit_flowchart_columns(ws, flow_range[0], flow_range[1])

    merged = {
        cell.coordinate
        for mr in ws.merged_cells.ranges
        for row in ws[mr.coord]
        for cell in row
    }
    for row in ws.iter_rows():
        if flow_range and flow_range[0] <= row[0].row <= flow_range[1]:
            continue
        for cell in row:
            if cell.coordinate in merged or cell.value is None:
                continue
            letter = get_column_letter(cell.column)
            w = min(44, max(10, len(str(cell.value)) * 1.1 + 2))
            ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, w)


def create_workbook() -> Workbook:
    reset_manifest()
    profiles = build_all_profiles()
    wb = Workbook()
    wb.remove(wb.active)

    build_index_sheet(wb.create_sheet("목차", 0), profiles)
    build_standard_sheets(wb)

    for domain_id, domain_name, screen_ids in DOMAINS:
        build_domain_sheet(wb.create_sheet(domain_id[:31]), domain_id, domain_name, screen_ids, profiles)

    for screen in SCREENS:
        sid = screen[0]
        ws = wb.create_sheet(sheet_title(sid, screen[1]))
        build_screen_process_sheet(ws, profiles[sid])

    for ws in wb.worksheets:
        _auto_fit_sheet(ws)

    return wb


def main() -> None:
    import subprocess

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    profiles = build_all_profiles()
    wb = create_workbook()
    wb.save(MASTER_FILE)
    save_manifest(MANIFEST_FILE)

    ps_script = ROOT / "scripts" / "add_excel_native_flowcharts.ps1"
    subprocess.run(
        [
            "powershell",
            "-NoProfile",
            "-ExecutionPolicy", "Bypass",
            "-File", str(ps_script),
            "-WorkbookPath", str(MASTER_FILE),
            "-ManifestPath", str(MANIFEST_FILE),
        ],
        check=True,
    )

    total_steps = sum(len(p.get("steps") or []) for p in profiles.values())
    total_nodes = sum(profile_summary(p)["흐름노드"] for p in profiles.values())
    print(f"Master: {MASTER_FILE}")
    print(f"  Sheets: {len(wb.sheetnames)} (목차 + 표준3 + 도메인{len(DOMAINS)} + 화면{len(SCREENS)})")
    print(f"  Flow nodes: {total_nodes}, Detail steps: {total_steps}")
    print("  Native Excel shapes: added by COM")


if __name__ == "__main__":
    main()
