# -*- coding: utf-8 -*-
"""단위 테스트 명세서 — 화면별 Excel 시트 생성."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from generate_screen_design_excel import (
    BORDER,
    DETAIL,
    HEADER_FILL,
    HEADER_FONT,
    SCREENS,
    SECTION_FILL,
    SECTION_FONT,
    WRAP,
    auto_fit_columns,
    sheet_title,
    write_data_table,
    write_kv_table,
    write_section_title,
)
from unit_test_cases import (
    COMMON_AUTH_CASES,
    COMMON_MSG_CASES,
    LAYER_BY_PATTERN,
    _backend_info,
    build_screen_test_cases,
    suggested_test_packages,
    test_summary,
)

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "unit-test" / "excel"
MASTER_FILE = OUT_DIR / "단위테스트_TO-BE.xlsx"

TC_HEADERS = [
    "TC ID",
    "유형",
    "테스트 대상",
    "시나리오",
    "전제조건",
    "입력/조건",
    "기대결과",
    "우선순위",
    "자동화",
    "상태",
]

PASS_FILL = PatternFill("solid", fgColor="DCFCE7")
P0_FILL = PatternFill("solid", fgColor="FEE2E2")
P1_FILL = PatternFill("solid", fgColor="FEF3C7")


def write_test_table(ws, row: int, cases: list[tuple]) -> int:
    for col, h in enumerate(TC_HEADERS, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row += 1
    for case in cases:
        padded = list(case[:8]) + ["JUnit 4", "미작성"]
        for col, val in enumerate(padded, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = BORDER
            c.alignment = WRAP
            if col == 8:
                if val == "P0":
                    c.fill = P0_FILL
                elif val == "P1":
                    c.fill = P1_FILL
        row += 1
    return row + 1


def build_screen_test_sheet(ws, screen: tuple) -> None:
    sid, name, domain, url, pattern, status, jsp, api, wireframe, fr, bpmn, notes = screen
    d = DETAIL.get(sid, {})
    backend = _backend_info(sid, url, jsp, d)
    cases = build_screen_test_cases(sid, screen, d)
    summary = test_summary(cases)
    p0 = sum(1 for c in cases if c[7] == "P0")

    row = 1
    ws.merge_cells("A1:J1")
    t = ws["A1"]
    t.value = f"단위테스트 명세 — SCR-{sid} {name}"
    t.font = Font(bold=True, size=14, color="166534")
    t.alignment = Alignment(horizontal="center", vertical="center")
    row = 3

    row = write_section_title(ws, row, "1. 기본 정보", 10)
    row = write_kv_table(ws, row, [
        ("화면ID", sid),
        ("화면명", name),
        ("도메인", domain),
        ("URL", url),
        ("패턴", pattern),
        ("상태", status),
        ("관련 FR", fr or "—"),
        ("API", api),
        ("총 TC", str(len(cases))),
        ("P0 TC", str(p0)),
    ])
    row += 1

    row = write_section_title(ws, row, "2. 테스트 범위", 10)
    layers = LAYER_BY_PATTERN.get(pattern, ["ApiController", "Service"])
    layer_text = ", ".join(layers)
    row = write_kv_table(ws, row, [
        ("대상 레이어", layer_text),
        ("View Controller", backend.get("controller", "—")),
        ("API Controller", backend.get("api", "—")),
        ("Service", backend.get("service", "—")),
        ("Mapper", backend.get("mapper", "—")),
        ("SQL XML", backend.get("sql", "—")),
        ("VO/DTO", backend.get("vo", "—")),
        ("권한", d.get("auth", "—")[:120]),
        ("유효성", d.get("validation", "—")[:120]),
    ])
    row += 1

    row = write_section_title(ws, row, "3. TC 요약", 10)
    summary_rows = [(k, str(v)) for k, v in sorted(summary.items())]
    row = write_data_table(ws, row, ["유형", "건수"], summary_rows)
    row += 1

    row = write_section_title(ws, row, "4. 테스트 클래스 (권장)", 10)
    pkg_rows = suggested_test_packages(sid, backend)
    row = write_data_table(ws, row, ["클래스", "경로"], pkg_rows)
    row += 1

    if d.get("tobe_changes"):
        row = write_section_title(ws, row, "5. TO-BE 변경 — 테스트 포인트", 10)
        for ch in d["tobe_changes"]:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            ws.cell(row=row, column=1, value=f"• {ch}").alignment = WRAP
            row += 1
        row += 1
        sec = 6
    else:
        sec = 5

    row = write_section_title(ws, row, f"{sec}. 테스트 케이스", 10)
    write_test_table(ws, row, cases)

    color = "166534" if status in ("신규", "변경", "v2") else "64748B"
    ws.sheet_properties.tabColor = color


def build_index_sheet(ws) -> None:
    ws.title = "목차"
    headers = [
        "No", "화면ID", "화면명", "도메인", "패턴", "상태",
        "TC수", "P0", "시트명", "화면설계 시트",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER

    widths = [5, 12, 22, 14, 8, 8, 8, 6, 22, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for i, s in enumerate(SCREENS, 1):
        sid, name, domain, url, pattern, status, jsp, *_rest = s
        d = DETAIL.get(sid, {})
        cases = build_screen_test_cases(sid, s, d)
        p0 = sum(1 for c in cases if c[7] == "P0")
        stitle = sheet_title(sid, name)
        design_stitle = stitle
        row = i + 1
        ws.cell(row=row, column=1, value=i).border = BORDER
        ws.cell(row=row, column=2, value=sid).border = BORDER
        ws.cell(row=row, column=3, value=name).border = BORDER
        ws.cell(row=row, column=4, value=domain).border = BORDER
        ws.cell(row=row, column=5, value=pattern).border = BORDER
        ws.cell(row=row, column=6, value=status).border = BORDER
        ws.cell(row=row, column=7, value=len(cases)).border = BORDER
        ws.cell(row=row, column=8, value=p0).border = BORDER
        c = ws.cell(row=row, column=9, value=stitle)
        c.border = BORDER
        c.hyperlink = f"#'{stitle}'!A1"
        c.font = Font(color="0563C1", underline="single")
        dc = ws.cell(row=row, column=10, value=design_stitle)
        dc.border = BORDER
        dc.hyperlink = f"../../screen-design/excel/화면설계서_TO-BE.xlsx#'{design_stitle}'!A1"
        dc.font = Font(color="0563C1", underline="single")

    total_tc = sum(
        len(build_screen_test_cases(s[0], s, DETAIL.get(s[0], {}))) for s in SCREENS
    )
    row = len(SCREENS) + 2
    ws.cell(row=row, column=1, value="합계").font = Font(bold=True)
    ws.cell(row=row, column=7, value=total_tc).font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "166534"


def build_standard_sheet(ws, title: str, rows: list[tuple[str, str]]) -> None:
    row = write_section_title(ws, 1, title, 4)
    write_data_table(ws, row, ["항목", "내용"], rows)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 56
    ws.sheet_properties.tabColor = "166534"


def build_common_tc_sheet(ws, title: str, cases: list[tuple]) -> None:
    row = write_section_title(ws, 1, title, 10)
    write_test_table(ws, row, cases)
    ws.sheet_properties.tabColor = "166534"


def create_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    build_index_sheet(wb.create_sheet("목차", 0))

    build_standard_sheet(wb.create_sheet("00_표준"), "JUnit 단위 테스트 표준", [
        ("프레임워크", "JUnit 4.12 + spring-test (MockMvc)"),
        ("패키지", "src/test/java/com/woori/ajs/{api,service,mapper}/"),
        ("네이밍", "클래스: {Target}Test / 메서드: test_{method}_{scenario}"),
        ("DB", "HSQLDB (로컬) — @Sql 또는 test fixture"),
        ("Mock", "Service 테스트: Mapper @Mock / API: Service @MockBean"),
        ("커버리지", "EMMA (pom.xml) — Service·ApiController 우선"),
        ("실행", "mvn test -Dtest={Class}Test"),
        ("우선순위", "P0=릴리스 차단, P1=스프린트 내, P2=여유 시"),
        ("상태", "미작성 → 작성중 → 완료 (수동 갱신)"),
        ("연계", "화면설계서 §8 API · §10 권한 · §11 유효성"),
    ])

    build_standard_sheet(wb.create_sheet("00_템플릿"), "테스트 메서드 템플릿 (Java)", [
        ("API 정상", "@Test public void test_list_ok() { mockMvc.perform(get(...)).andExpect(status().isOk()); }"),
        ("권한 거부", "@Test(expected=AuthException.class) public void test_list_no_access() { ... }"),
        ("Service", "@RunWith(MockitoJUnitRunner.class) @Mock Mapper; @InjectMocks Service;"),
        ("Mapper", "@Transactional @Sql(\"/testdata/{sid}_list.sql\") void test_selectList()"),
        ("Fixture", "src/test/resources/testdata/ — 화면ID별 SQL·JSON"),
    ])

    build_common_tc_sheet(wb.create_sheet("00_권한공통"), "공통 권한 테스트 케이스", COMMON_AUTH_CASES)
    build_common_tc_sheet(wb.create_sheet("00_메시지공통"), "공통 메시지·유효성 테스트", COMMON_MSG_CASES)

    for screen in SCREENS:
        sid, name = screen[0], screen[1]
        ws = wb.create_sheet(sheet_title(sid, name))
        build_screen_test_sheet(ws, screen)

    for ws in wb.worksheets:
        auto_fit_columns(ws, min_width=8, max_width=42)
        for col in range(1, 11):
            letter = get_column_letter(col)
            if (ws.column_dimensions[letter].width or 0) < 10:
                ws.column_dimensions[letter].width = 12

    return wb


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = create_workbook()
    wb.save(MASTER_FILE)
    total = sum(
        len(build_screen_test_cases(s[0], s, DETAIL.get(s[0], {}))) for s in SCREENS
    )
    print(f"Master: {MASTER_FILE}")
    print(f"  Sheets: {len(wb.sheetnames)} (목차 + 공통4 + 화면{len(SCREENS)})")
    print(f"  Test cases: {total} (+ common {len(COMMON_AUTH_CASES) + len(COMMON_MSG_CASES)})")


if __name__ == "__main__":
    main()
