# -*- coding: utf-8 -*-
"""화면설계서 — 화면별 Excel 시트 생성."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from excel_cell_mockup import MOCKUP_COLS
from screen_design_dev_data import BACKEND_MAP, PATTERN_IMPL
from screen_design_enrich import EVENTS_SPEC_NOTE, build_enriched_detail
from screen_design_tables import TABLE_CATALOG

# Excel 셀·서식 기반 레이아웃 목업 (HTML/PNG 와이어프레임 제외)
USE_EXCEL_CELL_MOCKUP = True

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "screen-design" / "excel"
MASTER_FILE = OUT_DIR / "화면설계서_TO-BE.xlsx"

# id, name, domain, url, pattern, status, jsp, api, wireframe, fr, bpmn, notes
SCREENS: list[tuple] = [
    ("LOGIN", "로그인", "00_공통", "/login", "단독", "v2",
     "login/login.jsp", "POST /login; /sso/prx", "login-v2.html", "FR-057",
     "로그인.bpmn", "SSO 연동"),
    ("MAIN", "메인프레임", "00_공통", "/index", "A", "v2",
     "layout/layout.jsp", "—", "main-shell.html", "FR-002, FR-003",
     "메인프레임.bpmn", "GNB+LNB+탭+iframe+SSE"),
    ("DASH", "대시보드_v1", "00_공통", "/dashboard", "B", "AS-IS",
     "dashboard/dashboard.jsp", "POST /api/dashboard", "dashboard.html", "—",
     "대시보드.bpmn", "v1 지도·차트"),
    ("DASH-v2", "대시보드_v2", "00_공통", "/dashboard", "B", "v2",
     "dashboard/dashboard.jsp", "POST /api/dashboard", "dashboard-v2.html",
     "FR-001, FR-003, FR-027", "대시보드.bpmn", "긴급 심사대상 카드"),
    ("9999", "사용자매뉴얼", "00_공통", "/common/manual", "C", "AS-IS",
     "common/manual/list.jsp", "GET /api/common/manual", "—", "—",
     "9999_사용자매뉴얼.bpmn", ""),
    ("GLOBAL", "전역UI_Toast_SSE", "00_공통", "—", "전역", "신규",
     "layout/layout.jsp", "GET /api/notify/stream; GET /api/notify/poll",
     "etc-fr005-urgent-notify.html", "FR-005", "—", "layout.jsp 단일 SSE"),
    ("1010", "본인작업_ToDoList", "01_심사", "/revert/todo", "C", "v2",
     "revert/todo/list.jsp", "GET /api/revert/todo", "list-1010-v2.html",
     "FR-004, FR-005", "1010_본인작업_ToDoList.bpmn", "S1 Task-Role + Depth"),
    ("1020", "진행업무별_현재상태", "01_심사", "/revert/status", "C", "AS-IS",
     "revert/status/list.jsp", "GET /api/revert/status", "—", "—",
     "1020_진행업무별_현재상태.bpmn", ""),
    ("1030", "업무미생성목록", "01_심사", "/revert/ungenerated", "C", "AS-IS",
     "revert/ungenerated/list.jsp",
     "GET /api/revert/ungenerated; POST confirm/cancel", "—", "—",
     "1030_업무미생성목록.bpmn", ""),
    ("1011", "심사상세", "01_심사", "POST /revert/detail", "D", "변경",
     "revert/detail/detail.jsp",
     "POST /api/common/detail/load; santionSave; santionApprv",
     "detail-1011-v2.html", "FR-006~029, FR-038, FR-064",
     "1011_심사상세.bpmn", "btnFlag=A; Depth 버튼 가변"),
    ("1011-BL", "BL번호수정", "01_심사", "모달", "G", "AS-IS",
     "revert/detail/ModiBlNum.jsp", "POST /api/common/revertDetail/ModiBlNum",
     "—", "—", "1011_BL번호수정.bpmn", ""),
    ("2010", "결재자_ToDoList", "02_결재", "/app/todo", "C", "변경",
     "app/todo/list.jsp", "GET /api/app/todo", "list-2010.html", "—",
     "2010_결재자_ToDoList.bpmn", "Depth=2만 표시"),
    ("2020", "결재자_진행상태", "02_결재", "/app/status", "C", "AS-IS",
     "app/status/list.jsp", "GET /api/app/status", "—", "—",
     "2020_결재자_진행상태.bpmn", ""),
    ("2011", "결재자_심사상세", "02_결재", "/app/detail/revert", "D", "AS-IS",
     "app/detail/detail.jsp", "POST /api/common/appDetail/*", "—", "—",
     "2011_결재자_심사상세.bpmn", "btnFlag=B; SoD"),
    ("2011-QA", "결재자_QA상세", "02_결재", "/app/detail/qa", "D", "AS-IS",
     "app/detail/qa/detail.jsp", "POST /api/common/appDetail/*", "—", "—", "—", "btnFlag=D"),
    ("BUNDLE", "일괄승인_팝업", "02_결재", "모달", "G", "AS-IS",
     "app/todo/bundle.jsp", "POST /api/app/todo/bundle", "—", "—",
     "2010_일괄승인팝업.bpmn", ""),
    ("3010", "QA_ToDoList", "03_QA", "/qa/todo", "C", "AS-IS",
     "qa/todo/list.jsp", "GET /api/qa/todo", "list-3010.html", "—",
     "3010_QA_ToDoList.bpmn", "QA Role 필터"),
    ("3020", "QA_진행상태", "03_QA", "/qa/status", "C", "AS-IS",
     "qa/status/list.jsp", "GET /api/qa/status", "—", "—", "3020_QA_진행상태.bpmn", ""),
    ("3011", "QA상세", "03_QA", "/qa/detail", "D", "AS-IS",
     "qa/detail/detail.jsp", "POST /api/common/qaDetail/*", "—", "—",
     "3011_QA상세.bpmn", "btnFlag=C"),
    ("4010", "업무생성목록", "04_현황명세", "/revert/stat/create", "C", "AS-IS",
     "revert/stat/create/list.jsp", "GET /api/revert/stat/create", "—", "—", "4010_*.bpmn", ""),
    ("4020", "담당자별_진행현황", "04_현황명세", "/revert/stat/progress", "C", "AS-IS",
     "revert/stat/progress/list.jsp", "GET /api/revert/stat/progress", "—", "—", "4020_*.bpmn", ""),
    ("4030", "심사완료명세", "04_현황명세", "/revert/stat/complete", "C", "AS-IS",
     "revert/stat/complete/list.jsp", "GET /api/revert/stat/complete", "—", "—", "4030_*.bpmn", ""),
    ("4040", "심사오류명세", "04_현황명세", "/revert/stat/error", "C", "AS-IS",
     "revert/stat/error/list.jsp", "GET /api/revert/stat/error", "—", "—", "4040_*.bpmn", ""),
    ("4050", "경보발생명세", "04_현황명세", "/revert/stat/alert", "C", "AS-IS",
     "revert/stat/alert/list.jsp", "GET /api/revert/stat/alert", "—", "—", "4050_*.bpmn", ""),
    ("5010", "업무마감", "05_업무일지", "/task/log/end", "F", "AS-IS",
     "task/log/end/list.jsp", "GET/POST /api/task/log/end", "—", "—", "5010_*.bpmn", ""),
    ("5020", "업무일지_등록", "05_업무일지", "/task/log/reg", "F", "AS-IS",
     "task/log/reg/list.jsp", "GET/POST /api/task/log/reg", "—", "—", "5020_*.bpmn", ""),
    ("5030", "업무일지_조회", "05_업무일지", "/task/log/inquiry", "C", "AS-IS",
     "task/log/inquiry/list.jsp", "GET /api/task/log/inquiry", "—", "—", "5030_*.bpmn", "차트"),
    ("5040", "성과관리", "05_업무일지", "/task/log/result", "C", "AS-IS",
     "task/log/result/list.jsp", "GET /api/task/log/result", "—", "—", "5040_*.bpmn", ""),
    ("6010", "성능분석", "06_통계", "/stat/analysis", "C", "AS-IS",
     "stat/analysis/list.jsp", "GET /api/stat/analysis", "—", "—", "6010_*.bpmn", ""),
    ("6020", "업무별_통계", "06_통계", "/stat/task", "C", "AS-IS",
     "stat/task/list.jsp", "GET /api/stat/task", "—", "—", "6020_*.bpmn", ""),
    ("6030", "항목별_통계", "06_통계", "/stat/item", "차트", "AS-IS",
     "stat/item/list.jsp", "GET /api/stat/item", "—", "—", "6030_*.bpmn", ""),
    ("6040", "조건별_통계", "06_통계", "/stat/cond", "C", "AS-IS",
     "stat/cond/list.jsp", "GET /api/stat/cond", "—", "—", "6040_*.bpmn", ""),
    ("6050", "고객별_통계", "06_통계", "/stat/user", "C", "AS-IS",
     "stat/user/list.jsp", "GET /api/stat/user", "—", "—", "6050_*.bpmn", ""),
    ("6066", "감독증빙_통계", "06_통계", "/stat/regulatory", "C", "신규",
     "stat/regulatory/list.jsp", "GET /api/stat/regulatory", "stats-6066-regulatory.html",
     "FR-066", "—", ""),
    ("6067", "업무현황_통계", "06_통계", "/stat/efficiency", "C", "신규",
     "stat/efficiency/list.jsp", "GET /api/stat/efficiency", "—", "FR-067", "—", ""),
    ("6069", "조기경보_통계", "06_통계", "/stat/early-warning", "C", "신규",
     "stat/early-warning/list.jsp", "GET /api/stat/early-warning", "—", "FR-069", "—", ""),
    ("7010", "사용자_관리", "07_관리자", "/admin/user", "E", "v2",
     "admin/user/list.jsp", "GET/POST /api/admin/user", "—", "FR-055",
     "7010_사용자_관리.bpmn", ""),
    ("7021", "로그인_이력", "07_관리자", "/admin/log/login", "C", "AS-IS",
     "admin/log/login/list.jsp", "GET /api/admin/log/login", "—", "—",
     "7021_로그인_이력관리.bpmn", ""),
    ("7022", "프로그램_사용이력", "07_관리자", "/admin/log/program", "C", "AS-IS",
     "admin/log/program/list.jsp", "GET /api/admin/log/program", "—", "—",
     "7022_프로그램_사용이력.bpmn", ""),
    ("7030", "권한별_메뉴관리", "07_관리자", "/admin/menu", "E", "AS-IS",
     "admin/menu/list.jsp", "GET/POST /api/admin/menu", "admin-7030.html", "—",
     "7030_권한별_메뉴관리.bpmn", ""),
    ("7031", "메뉴등록", "07_관리자", "/admin/menu/menuReg", "E", "AS-IS",
     "admin/menu/menuReg.jsp", "GET/POST /api/admin/menu/reg", "—", "—",
     "7031_메뉴등록.bpmn", "팝업"),
    ("7040", "WatchList_등록", "07_관리자", "/admin/watchlist", "E", "AS-IS",
     "admin/watchlist/watchInfo.jsp", "GET/POST /api/admin/watchlist/*",
     "admin-7040.html", "—", "7040_WatchList_등록.bpmn", ""),
    ("7050-RULE", "제재Rule_등록", "07_관리자", "/admin/sanction", "E", "AS-IS",
     "admin/sanction/sanctionInfo.jsp", "GET/POST /api/admin/sanction", "—", "—",
     "7050_제재Rule_등록.bpmn", "기존7050 번호"),
    ("7050", "업무별_권한설정", "07_관리자", "/admin/taskAuth", "E", "신규",
     "admin/taskAuth/list.jsp", "GET/POST /api/admin/taskAuth",
     "admin-7050-task-auth.html", "FR-042, FR-044, FR-054, FR-065", "—", "TO-BE 신규"),
    ("7051", "부점정보", "07_관리자", "/admin/branch", "C", "신규",
     "admin/branch/list.jsp", "GET/POST /api/admin/branch", "—", "FR-056", "—", ""),
    ("7052", "승인규칙_관리", "07_관리자", "/admin/approvDepth", "E", "신규",
     "admin/approvDepth/list.jsp",
     "GET/POST /api/admin/approvDepth; POST simulate",
     "admin-7052-approv-depth.html", "FR-041, FR-043", "—", ""),
    ("7060", "QA선정_관리", "07_관리자", "/admin/qa/target", "F", "AS-IS",
     "admin/qa/target.jsp", "GET/POST /api/admin/qa/target", "—", "—",
     "7060_QA선정_관리.bpmn", ""),
    ("7070", "부재_관리", "07_관리자", "/admin/absence", "C", "AS-IS",
     "admin/absence/list.jsp", "GET/POST /api/admin/absence", "—", "—",
     "7070_부재_관리.bpmn", ""),
    ("7080", "업무_재할당", "07_관리자", "/admin/reTask", "E", "변경",
     "admin/reTask/list.jsp", "GET/POST /api/admin/reTask", "—", "—",
     "7080_업무_재할당.bpmn", "SoD 연계"),
    ("7090", "일괄결재_의견관리", "07_관리자", "/admin/appMemo", "C", "AS-IS",
     "admin/appMemo/list.jsp", "GET/POST /api/admin/appMemo", "—", "—",
     "7090_일괄결재_의견관리.bpmn", ""),
    ("8010", "국가코드관리", "08_관리자시스템", "/admin/nation", "C", "AS-IS",
     "admin/nation/list.jsp", "GET/POST /api/admin/nation", "—", "—",
     "8010_국가코드관리.bpmn", ""),
    ("8020", "공통코드관리", "08_관리자시스템", "/admin/code", "E", "AS-IS",
     "admin/code/list.jsp", "GET/POST /api/admin/code", "—", "—",
     "8020_공통코드관리.bpmn", ""),
    ("8030", "영업일_관리", "08_관리자시스템", "/admin/businessDay", "F", "AS-IS",
     "admin/businessDay/list.jsp", "GET/POST /api/admin/businessDay", "—", "—",
     "8030_영업일_관리.bpmn", ""),
    ("8040", "도시항구관리", "08_관리자시스템", "/admin/city", "C", "AS-IS",
     "admin/city/list.jsp", "GET/POST /api/admin/city", "—", "—",
     "8040_도시항구관리.bpmn", ""),
    ("8050", "시스템심사진행현황", "08_관리자시스템", "/admin/status", "C", "AS-IS",
     "admin/status/list.jsp", "GET/POST /api/admin/status", "—", "—",
     "8050_시스템심사진행현황.bpmn", ""),
    ("8051", "재처리정보_등록", "08_관리자시스템", "팝업", "G", "AS-IS",
     "admin/status/reprocess.jsp", "POST /api/admin/status/reprocess", "—", "—",
     "8051_재처리정보_등록.bpmn", ""),
    ("8060", "후보정용어관리", "08_관리자시스템", "/admin/wordCorrection", "C", "AS-IS",
     "admin/wordCorrection/list.jsp", "GET/POST /api/admin/wordCorrection", "—", "—",
     "8060_후보정용어관리.bpmn", ""),
    ("8070", "AI_운영모니터링", "08_관리자시스템", "/admin/aiMonitor", "C", "신규",
     "admin/aiMonitor/list.jsp", "GET /api/admin/aiMonitor", "admin-8070-monitor.html",
     "FR-058, FR-064", "—", ""),
    ("9010", "심사상세_이력", "09_공통", "/common/history/detail", "D", "AS-IS",
     "common/history/detail.jsp", "POST /api/common/history/detail", "—", "—",
     "9010_심사상세_이력.bpmn", "btnFlag=E"),
    ("9020", "QA상세_이력", "09_공통", "/common/history/qa/detail", "D", "AS-IS",
     "common/history/qa/detail.jsp", "POST /api/common/history/qa/detail", "—", "—",
     "9020_QA상세_이력.bpmn", "btnFlag=F"),
    ("9080", "심사이력", "09_공통", "/common/revert/history", "H", "v2",
     "common/revert/history.jsp", "GET /api/common/revert/history", "history-9080.html",
     "FR-065", "9080_심사이력.bpmn", ""),
    ("9090", "QA이력", "09_공통", "/common/qa/history", "H", "AS-IS",
     "common/qa/history.jsp", "GET /api/common/qa/history", "—", "—",
     "9090_QA이력.bpmn", ""),
]

DETAIL = build_enriched_detail(SCREENS)

PATTERN_DESC = {
    "A": "메인 셸 — GNB+LNB+탭+iframe",
    "B": "대시보드 — 카드·긴급리스트",
    "C": "목록 — 검색+DataTables",
    "D": "상세팝업 — 3단분할",
    "E": "Master-Detail CRUD",
    "F": "폼/설정",
    "G": "모달",
    "H": "이력 타임라인",
    "단독": "로그인 단독",
    "전역": "전역 컴포넌트",
    "차트": "차트 통계",
}

# Styles
HEADER_FILL = PatternFill("solid", fgColor="1C5BBA")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
SECTION_FILL = PatternFill("solid", fgColor="EFF6FF")
SECTION_FONT = Font(bold=True, size=11, color="144A9C")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet_title(sid: str, name: str) -> str:
    """Excel sheet name max 31 chars."""
    title = f"{sid}_{name}" if len(f"{sid}_{name}") <= 31 else sid[:31]
    for ch in "[]:*?/\\":
        title = title.replace(ch, "_")
    return title


def write_section_title(ws, row: int, title: str, col_span: int = 6) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    cell = ws.cell(row=row, column=1, value=title)
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    cell.alignment = Alignment(vertical="center")
    for c in range(1, col_span + 1):
        ws.cell(row=row, column=c).border = BORDER
    return row + 1


def write_kv_table(ws, row: int, items: list[tuple[str, str]]) -> int:
    for label, value in items:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor="F8FAFC")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        ws.cell(row=row, column=2, value=value).alignment = WRAP
        for c in range(1, 7):
            ws.cell(row=row, column=c).border = BORDER
        row += 1
    return row


def write_data_table(ws, row: int, headers: list[str], rows: list[tuple]) -> int:
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    row += 1
    for data in rows:
        for col, val in enumerate(data, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = BORDER
            c.alignment = WRAP
        row += 1
    return row + 1


def display_width(value: Any) -> int:
    from excel_cell_mockup import display_width as mockup_display_width

    return mockup_display_width(value)


def auto_fit_columns(ws, *, min_width: float = 8, max_width: float = 48) -> None:
    """표 데이터 길이에 따라 열 너비를 조정한다.

    병합 셀의 긴 설명/제목은 제외한다.
    레이아웃 목업 영역 행은 제외하여 §4 Excel 셀 목업 전용 보정과 분리한다.
    """
    mockup_range = getattr(ws, "_mockup_range", None)
    merged_cells = {
        cell.coordinate
        for merged_range in ws.merged_cells.ranges
        for row in ws[merged_range.coord]
        for cell in row
    }
    max_by_col: dict[int, int] = {}

    for row in ws.iter_rows():
        if mockup_range and mockup_range[0] <= row[0].row <= mockup_range[1]:
            continue
        for cell in row:
            if cell.coordinate in merged_cells or cell.value is None:
                continue
            width = display_width(cell.value)
            if width:
                max_by_col[cell.column] = max(max_by_col.get(cell.column, 0), width)

    for col_idx, width in max_by_col.items():
        if mockup_range and col_idx <= MOCKUP_COLS:
            continue
        letter = get_column_letter(col_idx)
        current = ws.column_dimensions[letter].width or min_width
        adjusted = min(max_width, max(min_width, width * 1.15 + 2))
        ws.column_dimensions[letter].width = max(current, adjusted)


def infer_backend(url: str, jsp: str) -> dict[str, str]:
    if url in ("—", "", "모달", "팝업") or url.startswith("POST"):
        stem = Path(jsp).stem if jsp else "Screen"
        cap = stem[0].upper() + stem[1:] if stem else "Screen"
        return {"controller": f"*{cap}Controller", "api": f"*{cap}ApiController",
                "service": f"*{cap}Service", "mapper": f"*{cap}Mapper",
                "sql": f"*{cap}_SQL.xml", "vo": "SearchVO"}
    path = url.split()[-1] if " " in url else url
    parts = [p for p in path.strip("/").split("/") if p]
    if not parts:
        return {}
    prefix = {"admin": "Admin", "revert": "Revert", "app": "App", "qa": "Qa",
              "stat": "Stat", "task": "Task", "common": "Common"}.get(parts[0], parts[0].capitalize())
    name = parts[-1][0].upper() + parts[-1][1:]
    base = prefix + name
    return {"controller": f"{base}Controller", "api": f"{base}ApiController",
            "service": f"{base}Service / {base}ServiceImpl", "mapper": f"{base}Mapper",
            "sql": f"egovframework/sqlmap/ui/mappers/{base}_SQL.xml", "vo": "SearchVO, *VO"}


def build_checklist(sid: str, pattern: str, status: str, d: dict) -> list[str]:
    items = list(d.get("checklist") or [])
    items.extend(PATTERN_IMPL.get(pattern, {}).get("checklist") or [])
    if status == "신규":
        items.insert(0, "Controller / ApiController / Service / Mapper / JSP 신규 생성")
        items.append("메뉴(CSPD109TI) 및 screenAccess 등록")
    seen: set[str] = set()
    unique: list[str] = []
    for it in items:
        if it not in seen:
            seen.add(it)
            unique.append(it)
    return unique


def build_screen_sheet(ws, screen: tuple) -> None:
    sid, name, domain, url, pattern, status, jsp, api, wireframe, fr, bpmn, notes = screen
    d = DETAIL.get(sid, {})
    proc = d.get("process") or {}
    mockup_range: tuple[int, int] | None = None

    row = 1
    ws.merge_cells("A1:F1")
    t = ws["A1"]
    t.value = f"화면설계서 — SCR-{sid} {name}"
    t.font = Font(bold=True, size=14, color="1C5BBA")
    t.alignment = Alignment(horizontal="center", vertical="center")
    row = 3

    # 1. 기본 정보 (BPMN 제외)
    row = write_section_title(ws, row, "1. 기본 정보")
    row = write_kv_table(ws, row, [
        ("화면ID", sid),
        ("화면명", name),
        ("도메인", domain),
        ("URL", url),
        ("유형", f"{pattern} — {PATTERN_DESC.get(pattern, pattern)}"),
        ("상태", status),
        ("JSP", f"WEB-INF/jsp/{jsp}"),
        ("API", api),
        ("관련 FR", fr or "—"),
        ("프로세스 문서", proc.get("process_doc", "docs/process/목차.md")),
        ("비고", notes or "—"),
    ])
    row += 1

    # 2. 화면 목적
    purpose = d.get("purpose") or proc.get("purpose_process") or f"{name} 화면"
    row = write_section_title(ws, row, "2. 화면 목적")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=purpose).alignment = WRAP
    ws.row_dimensions[row].height = max(30, min(90, 15 * (1 + len(purpose) // 60)))
    row += 2

    # 3. 업무 프로세스
    row = write_section_title(ws, row, "3. 업무 프로세스")
    row = write_kv_table(ws, row, [
        ("선행 화면/조건", proc.get("nav_prev", "—")),
        ("후속 화면/결과", proc.get("nav_next", "—")),
        ("요약", proc.get("process_summary", purpose)),
    ])
    if proc.get("preconditions"):
        row = write_data_table(
            ws, row,
            ["구분", "조건"],
            [(a, b) for a, b in proc["preconditions"]],
        )
    steps = proc.get("process_steps") or []
    if steps:
        step_rows = []
        for s in steps:
            if len(s) >= 4:
                step_rows.append((s[0], s[1], s[2], s[3]))
            elif len(s) == 3:
                step_rows.append((s[0], s[1], s[2], ""))
            else:
                step_rows.append((s[0], s[1] if len(s) > 1 else "", "", ""))
        row = write_data_table(ws, row, ["단계", "처리", "트리거/설명", "후속·비고"], step_rows)
    if proc.get("status_codes"):
        row = write_data_table(
            ws, row,
            ["코드/조건", "의미"],
            list(proc["status_codes"]),
        )
    if proc.get("detail_doc"):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row=row, column=1, value=f"상세 프로세스: {proc['detail_doc']}").font = Font(
            italic=True, size=9, color="64748B"
        )
        row += 1
    row += 1

    # 4. 레이아웃 목업 (Excel 셀)
    row = write_section_title(ws, row, "4. 레이아웃 목업 (Excel 셀)")
    if USE_EXCEL_CELL_MOCKUP:
        from excel_cell_mockup import draw_excel_cell_mockup

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(
            row=row, column=1,
            value="검색·툴바·그리드·폼을 셀 병합/색상으로 표현 — Excel에서 직접 수정 가능 (HTML/PNG 목업 제외)",
        ).font = Font(italic=True, size=9, color="64748B")
        row += 1
        mockup_start = row
        row = draw_excel_cell_mockup(ws, row, sid, name, pattern, d, screen)
        mockup_range = (mockup_start, row - 1)
        row += 1

    sec = 5
    if d.get("search"):
        row = write_section_title(ws, row, f"{sec}. 검색 조건")
        row = write_data_table(ws, row, ["필드", "파라미터", "타입", "필수", "설명"], d["search"])
        sec += 1

    if d.get("fields"):
        row = write_section_title(ws, row, f"{sec}. 입력 필드")
        row = write_data_table(ws, row, ["필드", "타입", "필수", "DB/비고"], d["fields"])
        sec += 1

    if d.get("columns"):
        row = write_section_title(ws, row, f"{sec}. 목록·그리드 컬럼")
        cols = d["columns"]
        if cols and len(cols[0]) == 4:
            row = write_data_table(ws, row, ["컬럼", "data필드", "정렬", "비고"], cols)
        else:
            row = write_data_table(ws, row, ["컬럼", "정렬", "비고"], cols)
        sec += 1

    if d.get("buttons"):
        row = write_section_title(ws, row, f"{sec}. 버튼·액션")
        row = write_data_table(ws, row, ["조건", "버튼", "API/동작"], d["buttons"])
        sec += 1

    if d.get("events"):
        row = write_section_title(ws, row, f"{sec}. 이벤트·상호작용 (설계 설명)")
        note = d.get("events_note", EVENTS_SPEC_NOTE)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row=row, column=1, value=note).font = Font(italic=True, size=9, color="64748B")
        ws.cell(row=row, column=1).alignment = WRAP
        row += 1
        if d.get("interaction"):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            ws.cell(row=row, column=1, value=f"동작 요약: {d['interaction']}").alignment = WRAP
            row += 1
        row = write_data_table(ws, row, ["트리거", "동작", "비고"], d["events"])
        sec += 1

    # API
    row = write_section_title(ws, row, f"{sec}. API 명세")
    if d.get("api_specs"):
        rows = [(s[0], s[1], s[2] if len(s) > 2 else "", s[3] if len(s) > 3 else "", s[4] if len(s) > 4 else "")
                for s in d["api_specs"]]
        row = write_data_table(ws, row, ["경로", "Method", "Request", "Response", "비고"], rows)
    else:
        api_rows = []
        for part in api.split(";"):
            part = part.strip()
            if part and part != "—":
                if part.startswith(("GET", "POST", "PUT", "DELETE")):
                    m, _, p = part.partition(" ")
                    api_rows.append((p.strip() or part, m, "", "", ""))
                else:
                    api_rows.append((part, "—", "", "", ""))
        if not api_rows:
            api_rows = [("—", "—", "", "", "")]
        row = write_data_table(ws, row, ["경로", "Method", "Request", "Response", "비고"], api_rows)
    sec += 1

    row = write_section_title(ws, row, f"{sec}. 접근권한 (TO-BE)")
    auth = d.get("auth", "CSPD109TI + screenAccess + Task-Role (해당 시)")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row=row, column=1, value=auth).alignment = WRAP
    row += 2
    sec += 1

    row = write_section_title(ws, row, f"{sec}. 유효성·메시지")
    val = d.get("validation", "공통 세션·권한 검증 (AuthenticationInterceptor)")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    ws.cell(row=row, column=1, value=val).alignment = WRAP
    row += 2
    sec += 1

    be = d.get("backend") if isinstance(d.get("backend"), dict) and "controller" in d.get("backend", {}) else None
    if not be:
        be = BACKEND_MAP.get(sid) or infer_backend(url, jsp)
    row = write_section_title(ws, row, f"{sec}. 백엔드 구현")
    be_rows = []
    for key, label in [("controller", "View Controller"), ("api", "API Controller"),
                       ("service", "Service"), ("mapper", "Mapper"), ("sql", "SQL XML"), ("vo", "VO/DTO")]:
        if be.get(key) and be[key] != "—":
            be_rows.append((label, be[key], "src/.../com/woori/ajs/"))
    if be_rows:
        row = write_data_table(ws, row, ["구분", "클래스/파일", "경로"], be_rows)
    if d.get("flow"):
        row = write_section_title(ws, row, f"{sec}.1 처리 흐름 (구현)")
        for i, step in enumerate(d["flow"], 1):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            text = step.lstrip("0123456789. ")
            ws.cell(row=row, column=1, value=f"{i}. {text}").alignment = WRAP
            row += 1
        row += 1
    sec += 1

    if d.get("tables"):
        row = write_section_title(ws, row, f"{sec}. 관련 테이블")
        row = write_data_table(ws, row, ["논리명", "물리명", "용도"], d["tables"])
        sec += 1

    fe = d.get("frontend")
    if fe or jsp:
        row = write_section_title(ws, row, f"{sec}. 프론트엔드 구현")
        fe_rows = []
        if isinstance(fe, dict):
            for k, label in [("jsp", "JSP"), ("view", "View Name"), ("js", "JavaScript"),
                             ("css", "CSS"), ("includes", "Include"), ("forms", "Form")]:
                if fe.get(k) and "wireframe" not in k.lower() and "wireframes/" not in str(fe.get(k, "")).lower():
                    fe_rows.append((label, str(fe[k]), ""))
                elif fe.get(k) and k == "js" and "wireframes/" in str(fe[k]):
                    fe_rows.append((label, str(fe[k]).replace("docs/wireframes/js/", "js/"), "프론트 스크립트"))
        if not fe_rows:
            fe_rows.append(("JSP", f"WEB-INF/jsp/{jsp}", "blank/ 뷰"))
        for item in PATTERN_IMPL.get(pattern, {}).get("frontend", []):
            if len(item) >= 2 and not any(item[0] == r[0] for r in fe_rows):
                fe_rows.append((item[0], item[1], item[2] if len(item) > 2 else ""))
        row = write_data_table(ws, row, ["구분", "파일/함수", "설명"], fe_rows)
        sec += 1

    row = write_section_title(ws, row, f"{sec}. 화면 상태")
    row = write_data_table(ws, row, ["상태", "설명"], [
        ("초기", "View Controller 모델 로드, 그리드 empty"),
        ("조회중", "initLoadingDisplay 로딩 표시"),
        ("목록/상세표시", "DataTables·폼 바인딩 / rowCallback 스타일"),
        ("상세오픈", "팝업 fullLayout, 부모 유지"),
        ("갱신", "SSE / 팝업닫기 → tableReload"),
    ])
    sec += 1

    if d.get("tobe_changes"):
        row = write_section_title(ws, row, f"{sec}. TO-BE 변경")
        for ch in d["tobe_changes"]:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            ws.cell(row=row, column=1, value=f"• {ch}").alignment = WRAP
            row += 1
        sec += 1

    checklist = build_checklist(sid, pattern, status, d)
    if checklist:
        row = write_section_title(ws, row, f"{sec}. 개발 체크리스트")
        for it in checklist:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            ws.cell(row=row, column=1, value=f"☐ {it}").alignment = WRAP
            row += 1
        sec += 1

    row = write_section_title(ws, row, f"{sec}. 관련 문서")
    doc_rows = [
        ("프로세스 정의서", proc.get("process_doc", "docs/process/목차.md")),
    ]
    if proc.get("detail_doc"):
        doc_rows.append(("상세 프로세스", proc["detail_doc"]))
    doc_rows.extend([
        ("UI 표준", "docs/screen-design/00_UI표준.md"),
        ("권한·API", "docs/screen-design/00_권한_API_매트릭스.md"),
        ("메시지", "docs/screen-design/00_메시지_검증.md"),
        ("시스템 프로세스", "docs/system-process.md"),
    ])
    row = write_data_table(ws, row, ["문서", "경로"], doc_rows)

    ws.sheet_properties.tabColor = "1C5BBA" if status in ("신규", "변경", "v2") else "94A3B8"
    if mockup_range:
        ws._mockup_range = mockup_range  # type: ignore[attr-defined]


def build_index_sheet(ws) -> None:
    ws.title = "목차"
    headers = ["No", "화면ID", "화면명", "도메인", "유형", "상태", "프로세스문서", "시트명"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
    widths = [5, 12, 22, 14, 8, 8, 36, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for i, s in enumerate(SCREENS, 1):
        sid, name, domain, url, pattern, status, *_rest = s
        stitle = sheet_title(sid, name)
        d = DETAIL.get(sid, {})
        proc_path = (d.get("process") or {}).get("process_doc", "docs/process/목차.md")
        row = i + 1
        ws.cell(row=row, column=1, value=i).border = BORDER
        ws.cell(row=row, column=2, value=sid).border = BORDER
        ws.cell(row=row, column=3, value=name).border = BORDER
        ws.cell(row=row, column=4, value=domain).border = BORDER
        ws.cell(row=row, column=5, value=pattern).border = BORDER
        ws.cell(row=row, column=6, value=status).border = BORDER
        ws.cell(row=row, column=7, value=proc_path).border = BORDER
        c = ws.cell(row=row, column=8, value=stitle)
        c.border = BORDER
        c.hyperlink = f"#'{stitle}'!A1"
        c.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "144A9C"


def build_standard_sheet(ws, section_title: str, rows: list[tuple[str, str]]) -> None:
    row = write_section_title(ws, 1, section_title, 4)
    write_data_table(ws, row, ["항목", "내용"], rows)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 50


def create_workbook(include_index: bool = True) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    if include_index:
        build_index_sheet(wb.create_sheet("목차", 0))

    build_standard_sheet(wb.create_sheet("00_프로세스"), "업무 프로세스 개요", [
        ("E2E 흐름", "WINI접수 → 업무생성 → 자동심사 → 1010/1011 → 2010/2011 → (QA) → 완료명세/통계"),
        ("프로세스 정의서", "docs/process/*.md — 도메인별 화면 프로세스"),
        ("상세 프로세스", "docs/process/상세/*.md — 단계·API·분기·예외"),
        ("시스템 프로세스", "docs/system-process.md — End-to-End·상태전이"),
        ("btnFlag", "A=심사 B=결재 C=QA D=QA결재 E/F=이력읽기전용"),
        ("상태코드", "AI_INPT_PROS_CD / AI_INPT_APPV_STCD — 화면별 조회 조건"),
        ("공통패턴", "목록: 검색→API→DataTables→상세 · 상세: load→편집→저장/상신/승인"),
    ])
    build_standard_sheet(wb.create_sheet("00_UI표준"), "UI/UX 설계 표준", [
        ("Primary", "#1c5bba — GNB, 버튼"),
        ("폰트", "Noto Sans KR 13px"),
        ("목록패턴", "검색영역 + DataTables + 서버 페이징"),
        ("상세패턴", "3단분할 fullLayout + btnFlag"),
        ("결과30", "빨강 Alert"), ("결과40", "녹색 Clean"),
    ])
    build_standard_sheet(wb.create_sheet("00_권한API"), "권한·API 매트릭스", [
        ("1010", "S1 Task-Role + screenAccess"),
        ("2010", "S2 + Depth=2"),
        ("7050/7052", "mgpeYn=Y"),
        ("메뉴조건", "CSPD109TI AND screenAccess AND Task-Role"),
    ])
    build_standard_sheet(wb.create_sheet("00_메시지"), "유효성·메시지", [
        ("MSG-SOD-001", "동일 업무 S1·S2 동시 부여 불가"),
        ("MSG-SOD-002", "본인 처리 건 결재 불가"),
        ("MSG-DEPTH-001", "단독결재 최종승인 안내"),
    ])
    catalog_rows = [
        (logical, physical, "sql/01_ddl.sql" if physical <= "CSPD900TI" else "TO-BE 신규")
        for physical, logical in sorted(TABLE_CATALOG.items())
    ]
    ws_tbl = wb.create_sheet("00_테이블")
    row = write_section_title(ws_tbl, 1, "테이블 정의 (논리명 / 물리명)", 4)
    write_data_table(ws_tbl, row, ["논리명", "물리명", "비고"], catalog_rows)
    ws_tbl.column_dimensions["A"].width = 28
    ws_tbl.column_dimensions["B"].width = 14
    ws_tbl.column_dimensions["C"].width = 18

    for screen in SCREENS:
        sid, name = screen[0], screen[1]
        ws = wb.create_sheet(sheet_title(sid, name))
        build_screen_sheet(ws, screen)

    for ws in wb.worksheets:
        auto_fit_columns(ws)
        mockup_range = getattr(ws, "_mockup_range", None)
        if mockup_range:
            from excel_cell_mockup import auto_fit_mockup_columns

            auto_fit_mockup_columns(ws, mockup_range[0], mockup_range[1])

    return wb


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    master = create_workbook(include_index=True)
    master.save(MASTER_FILE)
    print(f"Master: {MASTER_FILE} ({len(master.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
