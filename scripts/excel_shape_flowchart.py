# -*- coding: utf-8 -*-
"""Pillow 기반 플로우차트 — 도형(타원·사각·마름모) + 화살표 PNG → Excel 삽입."""
from __future__ import annotations

import math
import tempfile
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image, ImageDraw, ImageFont

C_START = (34, 197, 94)
C_END = (100, 116, 139)
C_TASK_FILL = (237, 233, 254)
C_TASK_BORDER = (124, 58, 237)
C_GATE_FILL = (254, 240, 138)
C_GATE_BORDER = (202, 138, 4)
C_ARROW = (124, 58, 237)
C_LOOP = (194, 65, 12)
C_BG = (255, 255, 255)
C_TEXT = (30, 41, 59)
C_WHITE = (255, 255, 255)

FLOW_COLS = 16
ANCHOR_COL = 2


@dataclass
class NodeLayout:
    nid: str
    ntype: str
    name: str
    x: int
    y: int
    w: int
    h: int


@dataclass
class EdgeLayout:
    x1: int
    y1: int
    x2: int
    y2: int
    label: str = ""
    loop: bool = False
    elbow: bool = False


@dataclass
class FlowLayout:
    nodes: list[NodeLayout] = field(default_factory=list)
    edges: list[EdgeLayout] = field(default_factory=list)
    width: int = 800
    height: int = 400


def _load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for path in (
        [r"C:\Windows\Fonts\malgunbd.ttf", r"C:\Windows\Fonts\malgun.ttf"]
        if bold
        else [r"C:\Windows\Fonts\malgun.ttf", r"C:\Windows\Fonts\arial.ttf"]
    ):
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                pass
    return ImageFont.load_default()


def _steps_to_flow(steps: list) -> tuple[list, list]:
    elements = [("Start", "시작", "시작", "startEvent")]
    flows: list[tuple] = []
    prev = "Start"
    for i, step in enumerate(steps[:10], 1):
        name = str(step[1] if len(step) > 1 else f"단계{i}")[:24]
        eid = f"S{i}"
        elements.append((eid, "업무", name, "task"))
        flows.append((f"F{i}", prev, eid, "", ""))
        prev = eid
    elements.append(("End", "종료", "종료", "endEvent"))
    flows.append(("F_end", prev, "End", "", ""))
    return elements, flows


def _build_graph(bpmn_data: dict[str, Any]):
    elements = bpmn_data.get("elements") or []
    flows = bpmn_data.get("flows") or []
    elem_map = {e[0]: e for e in elements}
    outgoing: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for _, src, tgt, label, _ in flows:
        outgoing[src].append((tgt, label or ""))
    start_id = next((e[0] for e in elements if e[1] == "시작"), None)
    return elem_map, outgoing, start_id


def _node_size(etype: str) -> tuple[int, int]:
    if etype in ("시작", "종료"):
        return 80, 40
    if etype == "분기":
        return 108, 58
    return 156, 46


def _layout_flowchart(elem_map: dict, outgoing: dict, start_id: str | None) -> FlowLayout:
    layout = FlowLayout()
    if not start_id or start_id not in elem_map:
        return layout

    nodes: dict[str, NodeLayout] = {}
    edges: list[EdgeLayout] = []
    visited: set[str] = set()
    cx_main = 380
    y = 36
    v_gap = 24

    def add_edge(src: str, tgt: str, label: str = "", *, loop: bool = False, elbow: bool = False) -> None:
        if src not in nodes or tgt not in nodes:
            return
        s, t = nodes[src], nodes[tgt]
        edges.append(EdgeLayout(
            s.x + s.w // 2, s.y + s.h,
            t.x + t.w // 2, t.y,
            label=label, loop=loop, elbow=elbow,
        ))

    def place(nid: str, x: int, py: int) -> NodeLayout:
        _, etype, name, _ = elem_map[nid]
        w, h = _node_size(etype)
        n = NodeLayout(nid, etype, name, x - w // 2, py, w, h)
        nodes[nid] = n
        layout.nodes.append(n)
        return n

    def bottom(n: NodeLayout) -> int:
        return n.y + n.h

    def walk(nid: str, x: int, py: int, depth: int = 0) -> int:
        if nid in visited or depth > 24 or nid not in elem_map:
            return py
        visited.add(nid)
        n = place(nid, x, py)
        py = bottom(n) + v_gap
        children = outgoing.get(nid, [])
        etype = elem_map[nid][1]

        if etype == "분기" and len(children) > 1:
            branch_y = py + 8
            n_branches = len(children)
            span = max(480, (n_branches - 1) * 180)
            xs = [x - span // 2 + i * (span // max(n_branches - 1, 1)) for i in range(n_branches)]
            max_bottom = branch_y

            for i, (child_id, label) in enumerate(children):
                bx = xs[i]
                if child_id in visited:
                    add_edge(nid, child_id, label or "↺", loop=True, elbow=True)
                    continue
                add_edge(nid, child_id, label)

                child_etype = elem_map.get(child_id, ("", "업무", child_id, ""))[1]
                if child_etype == "종료":
                    cn = place(child_id, bx, branch_y)
                    visited.add(child_id)
                    max_bottom = max(max_bottom, bottom(cn))
                    continue

                cn = place(child_id, bx, branch_y)
                visited.add(child_id)
                cb = bottom(cn)
                max_bottom = max(max_bottom, cb)

                for tgt2, lbl2 in outgoing.get(child_id, []):
                    if tgt2 in nodes and tgt2 != child_id:
                        add_edge(child_id, tgt2, lbl2 or "↺", loop=True, elbow=True)
                    elif tgt2 in elem_map and elem_map[tgt2][1] == "종료" and tgt2 not in visited:
                        en = place(tgt2, bx, cb + v_gap)
                        visited.add(tgt2)
                        add_edge(child_id, tgt2)
                        max_bottom = max(max_bottom, bottom(en))

            return max_bottom + v_gap + 20

        for child_id, label in children:
            if child_id in visited:
                if child_id in nodes:
                    add_edge(nid, child_id, label or "↺", loop=True, elbow=True)
                continue
            add_edge(nid, child_id, label)
            if elem_map.get(child_id, ("", "", "", ""))[1] == "종료":
                place(child_id, x, py)
                visited.add(child_id)
                return bottom(nodes[child_id]) + v_gap
            py = walk(child_id, x, py, depth + 1)
        return py

    walk(start_id, cx_main, y)
    layout.edges = edges

    if layout.nodes:
        min_x = min(n.x for n in layout.nodes) - 36
        max_x = max(n.x + n.w for n in layout.nodes) + 36
        max_y = max(n.y + n.h for n in layout.nodes) + 44
        if min_x < 20:
            shift = 20 - min_x
            for n in layout.nodes:
                n.x += shift
            for e in layout.edges:
                e.x1 += shift
                e.x2 += shift
            max_x += shift
        layout.width = max(560, max_x)
        layout.height = max(240, max_y)
    return layout


def _wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_w: int) -> list[str]:
    words = text.replace("\n", " ").split()
    if not words:
        return [text[:14]]
    lines, cur = [], words[0]
    for w in words[1:]:
        test = f"{cur} {w}"
        if draw.textlength(test, font=font) <= max_w:
            cur = test
        else:
            lines.append(cur)
            cur = w
    lines.append(cur)
    return lines[:3]


def _draw_arrow(draw: ImageDraw.ImageDraw, edge: EdgeLayout) -> None:
    color = C_LOOP if edge.loop else C_ARROW
    x1, y1, x2, y2 = edge.x1, edge.y1, edge.x2, edge.y2
    if edge.elbow and abs(x2 - x1) > 20:
        mid_y = (y1 + y2) // 2
        draw.line((x1, y1, x1, mid_y), fill=color, width=2)
        draw.line((x1, mid_y, x2, mid_y), fill=color, width=2)
        draw.line((x2, mid_y, x2, y2), fill=color, width=2)
        ax, ay = x2, y2
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=2)
        ax, ay = x2, y2
    if abs(ax - x1) < 4 and abs(ay - y1) < 4:
        return
    angle = math.atan2(ay - (y1 if not edge.elbow else ay), ax - x1)
    size = 9
    pts = [
        (ax, ay),
        (ax - size * math.cos(angle - 0.45), ay - size * math.sin(angle - 0.45)),
        (ax - size * math.cos(angle + 0.45), ay - size * math.sin(angle + 0.45)),
    ]
    draw.polygon(pts, fill=color)
    if edge.label and not edge.loop:
        mx, my = (x1 + x2) // 2 - 20, (y1 + y2) // 2 - 14
        draw.text((mx, my), edge.label[:12], fill=(100, 116, 139), font=_load_font(8))


def _render_png(layout: FlowLayout, title: str, out_path: Path) -> None:
    img = Image.new("RGB", (layout.width, layout.height), C_BG)
    draw = ImageDraw.Draw(img)
    font = _load_font(11)
    font_s = _load_font(9)
    font_title = _load_font(13, bold=True)

    draw.text((14, 8), title, fill=C_TASK_BORDER, font=font_title)
    draw.text((14, layout.height - 20),
              "●시작(타원)  ■종료(타원)  ◆분기(마름모)  □업무(사각)  →화살표  ↺되돌아감",
              fill=(100, 116, 139), font=font_s)

    for edge in layout.edges:
        _draw_arrow(draw, edge)

    for node in layout.nodes:
        x, y, w, h = node.x, node.y, node.w, node.h
        if node.ntype == "시작":
            draw.ellipse((x, y, x + w, y + h), fill=C_START, outline=(22, 163, 74), width=2)
            tw = draw.textlength("시작", font=font)
            draw.text((x + (w - tw) / 2, y + 11), "시작", fill=C_WHITE, font=font)
        elif node.ntype == "종료":
            draw.ellipse((x, y, x + w, y + h), fill=C_END, outline=(71, 85, 105), width=2)
            tw = draw.textlength("종료", font=font)
            draw.text((x + (w - tw) / 2, y + 11), "종료", fill=C_WHITE, font=font)
        elif node.ntype == "분기":
            cx, cy = x + w // 2, y + h // 2
            draw.polygon([(cx, y + 2), (x + w - 2, cy), (cx, y + h - 2), (x + 2, cy)],
                         fill=C_GATE_FILL, outline=C_GATE_BORDER)
            for i, line in enumerate(_wrap_text(draw, node.name, font_s, w - 20)):
                tw = draw.textlength(line, font=font_s)
                draw.text((cx - tw / 2, cy - 8 + i * 12), line, fill=C_TEXT, font=font_s)
        else:
            draw.rounded_rectangle((x, y, x + w, y + h), radius=10, fill=C_TASK_FILL,
                                   outline=C_TASK_BORDER, width=2)
            lines = _wrap_text(draw, node.name, font, w - 14)
            ty = y + (h - len(lines) * 15) // 2
            for line in lines:
                tw = draw.textlength(line, font=font)
                draw.text((x + (w - tw) / 2, ty), line, fill=C_TEXT, font=font)
                ty += 15

    out_path.parent.mkdir(parents=True, exist_ok=True)
    img.save(out_path, format="PNG")


def render_flowchart(
    ws: Worksheet,
    start_row: int,
    bpmn_data: dict[str, Any],
    *,
    steps_fallback: list | None = None,
    title: str = "",
    screen_id: str = "",
    cache_dir: Path | None = None,
) -> tuple[int, tuple[int, int]]:
    elements = bpmn_data.get("elements") or []
    flows = bpmn_data.get("flows") or []
    if not elements and steps_fallback:
        elements, flows = _steps_to_flow(steps_fallback)

    elem_map, outgoing, start_id = _build_graph({"elements": elements, "flows": flows})
    flow_start = start_row

    if not elem_map or not start_id:
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=10)
        ws.cell(row=start_row, column=1, value="흐름 정의 없음")
        return start_row + 2, (flow_start, start_row + 1)

    layout = _layout_flowchart(elem_map, outgoing, start_id)
    chart_title = title or "업무 흐름도"
    if cache_dir is None:
        cache_dir = Path(tempfile.gettempdir()) / "ta_ui_flowcharts"
    png_path = cache_dir / f"{screen_id or 'flow'}.png"
    _render_png(layout, chart_title, png_path)

    img = XLImage(str(png_path))
    scale = min(1.0, 680 / max(layout.width, 1))
    img.width = int(layout.width * scale)
    img.height = int(layout.height * scale)
    ws.add_image(img, f"{get_column_letter(ANCHOR_COL + 1)}{start_row}")

    rows_needed = max(14, int(img.height / 16) + 2)
    for r in range(start_row, start_row + rows_needed):
        ws.row_dimensions[r].height = 16

    flow_end = start_row + rows_needed
    return flow_end + 1, (flow_start, flow_end)


def fit_flowchart_columns(ws: Worksheet, start_row: int, end_row: int) -> None:
    for c in range(1, FLOW_COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10
