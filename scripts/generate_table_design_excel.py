# -*- coding: utf-8 -*-
"""ERD(sanction.erd) → 테이블 정의서 Excel 생성."""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
ERD_FILE = ROOT / "sql" / "sanction.erd"
OUT_DIR = ROOT / "docs" / "table-design" / "excel"
OUT_FILE = OUT_DIR / "테이블정의서.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1C5BBA")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1C5BBA")
META_FONT = Font(size=10, color="475569")
TABLE_FILL = PatternFill("solid", fgColor="E8F1FC")
TABLE_FONT = Font(bold=True, color="1C5BBA")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

COLUMNS = [
    "테이블명",
    "테이블설명",
    "No",
    "컬럼명",
    "데이터타입",
    "NULL허용",
    "PK",
    "FK",
    "컬럼설명",
    "출처",
]


def to_excel_type(data_type: str) -> str:
    """Tibero/Oracle 타입 → 테이블정의서 표기."""
    if not data_type:
        return "varchar(100)"
    d = data_type.strip()
    upper = d.upper()
    m = re.match(r"VARCHAR2\((\d+)\)", upper)
    if m:
        return f"varchar({m.group(1)})"
    m = re.match(r"CHAR\((\d+)\)", upper)
    if m:
        return f"char({m.group(1)})"
    m = re.match(r"NUMBER\((\d+),(\d+)\)", upper)
    if m:
        return f"numeric({m.group(1)},{m.group(2)})"
    if upper == "NUMBER(19)":
        return "bigint"
    if upper.startswith("NUMBER"):
        return "numeric(19)"
    if upper == "CLOB":
        return "text"
    return d.lower()


def load_erd(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def col_name_from_id(col_id: str, columns: dict) -> str | None:
    col = columns.get(col_id)
    if col:
        return col["name"]
    # ERD 관계에만 존재하는 컬럼 ID fallback (col_{TABLE}_{COL})
    m = re.match(r"col_[A-Z0-9]+_(.+)", col_id)
    return m.group(1) if m else None


def build_fk_map(erd: dict) -> dict[tuple[str, str], str]:
    """(table, column) → 참조 테이블.컬럼"""
    tables = erd["collections"]["tableEntities"]
    columns = erd["collections"]["tableColumnEntities"]
    fk_map: dict[tuple[str, str], str] = {}

    for rel in erd["collections"].get("relationshipEntities", {}).values():
        start = rel["start"]
        end = rel["end"]
        start_table = tables.get(start["tableId"], {}).get("name")
        end_table = tables.get(end["tableId"], {}).get("name")
        if not start_table or not end_table or not start.get("columnIds") or not end.get("columnIds"):
            continue
        start_col = col_name_from_id(start["columnIds"][0], columns)
        end_col = col_name_from_id(end["columnIds"][0], columns)
        if not start_col or not end_col:
            continue
        fk_map[(end_table, end_col)] = f"{start_table}.{start_col}"

    return fk_map


def parse_tables(erd: dict) -> list[dict]:
    table_entities = erd["collections"]["tableEntities"]
    column_entities = erd["collections"]["tableColumnEntities"]
    fk_map = build_fk_map(erd)
    table_ids = erd["doc"]["tableIds"]

    rows: list[dict] = []
    for tid in table_ids:
        table = table_entities[tid]
        tname = table["name"]
        tcomment = table.get("comment", "")

        for seq, cid in enumerate(table.get("seqColumnIds") or table.get("columnIds", []), 1):
            col = column_entities[cid]
            cname = col["name"]
            is_pk = bool(col.get("options", 0) & 2)
            fk_ref = fk_map.get((tname, cname), "")
            is_fk = bool(fk_ref) or bool(col.get("ui", {}).get("keys", 0) & 2)

            rows.append(
                {
                    "table": tname,
                    "table_comment": tcomment,
                    "seq": seq,
                    "column": cname,
                    "type": to_excel_type(col.get("dataType", "")),
                    "nullable": "NO" if is_pk else "YES",
                    "pk": "Y" if is_pk else "",
                    "fk": fk_ref if is_fk else "",
                    "comment": col.get("comment", ""),
                    "source": "sql/sanction.erd",
                }
            )
    return rows


def style_header_row(ws, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER


def write_meta(ws) -> int:
    db_name = "AI_SANCTION"
    erd = load_erd(ERD_FILE)
    db_name = erd.get("settings", {}).get("databaseName", db_name)
    table_count = len(erd["doc"]["tableIds"])
    col_count = len(erd["collections"]["tableColumnEntities"])

    ws.merge_cells("A1:J1")
    title = ws.cell(row=1, column=1, value="AI 심사 시스템 테이블 정의서")
    title.font = TITLE_FONT
    title.alignment = Alignment(vertical="center")

    meta_lines = [
        f"데이터베이스: {db_name}",
        f"생성일: {date.today().isoformat()}",
        f"출처: sql/sanction.erd",
        f"테이블 {table_count}개 · 컬럼 {col_count}개",
    ]
    for i, line in enumerate(meta_lines, 2):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=10)
        cell = ws.cell(row=i, column=1, value=line)
        cell.font = META_FONT
    return 6


def auto_fit_columns(ws, min_width: int = 8, max_width: int = 42) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = min_width
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is None:
                    continue
                max_len = max(max_len, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = max_len


def build_definition_sheet(ws, rows: list[dict]) -> None:
    ws.title = "테이블정의서"
    header_row = write_meta(ws)
    style_header_row(ws, header_row, COLUMNS)

    row_idx = header_row + 1
    prev_table = None
    for item in rows:
        if item["table"] != prev_table:
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=col).fill = TABLE_FILL
            prev_table = item["table"]

        values = [
            item["table"],
            item["table_comment"],
            item["seq"],
            item["column"],
            item["type"],
            item["nullable"],
            item["pk"],
            item["fk"],
            item["comment"],
            item["source"],
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = BORDER
            cell.alignment = WRAP if col in (2, 9, 10) else CENTER
            if col in (3, 6, 7):
                cell.alignment = CENTER
        row_idx += 1

    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_properties.tabColor = "1C5BBA"
    widths = [14, 22, 5, 28, 16, 8, 5, 24, 32, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_index_sheet(ws, rows: list[dict]) -> None:
    ws.title = "목차"
    headers = ["No", "테이블명", "테이블설명", "컬럼수", "PK", "시트"]
    style_header_row(ws, 1, headers)

    table_info: dict[str, dict] = {}
    for item in rows:
        tname = item["table"]
        if tname not in table_info:
            table_info[tname] = {
                "comment": item["table_comment"],
                "cols": 0,
                "pk": [],
            }
        table_info[tname]["cols"] += 1
        if item["pk"] == "Y":
            table_info[tname]["pk"].append(item["column"])

    for i, (tname, info) in enumerate(table_info.items(), 1):
        row = i + 1
        pk_str = ", ".join(info["pk"])
        ws.cell(row=row, column=1, value=i).border = BORDER
        ws.cell(row=row, column=2, value=tname).border = BORDER
        ws.cell(row=row, column=3, value=info["comment"]).border = BORDER
        ws.cell(row=row, column=4, value=info["cols"]).border = BORDER
        ws.cell(row=row, column=5, value=pk_str).border = BORDER
        link = ws.cell(row=row, column=6, value="테이블정의서")
        link.border = BORDER
        link.hyperlink = f"#'테이블정의서'!A1"
        link.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "144A9C"
    for i, w in enumerate([5, 14, 24, 8, 28, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_fk_sheet(ws, erd: dict) -> None:
    ws.title = "FK관계"
    headers = ["No", "자식테이블", "자식컬럼", "부모테이블", "부모컬럼"]
    style_header_row(ws, 1, headers)

    tables = erd["collections"]["tableEntities"]
    columns = erd["collections"]["tableColumnEntities"]
    row = 2
    rel_no = 0
    for rel in erd["collections"].get("relationshipEntities", {}).values():
        start = rel["start"]
        end = rel["end"]
        parent_table = tables.get(start["tableId"], {}).get("name")
        child_table = tables.get(end["tableId"], {}).get("name")
        if not parent_table or not child_table:
            continue
        parent_col = col_name_from_id(start["columnIds"][0], columns) if start.get("columnIds") else None
        child_col = col_name_from_id(end["columnIds"][0], columns) if end.get("columnIds") else None
        if not parent_col or not child_col:
            continue
        rel_no += 1
        values = [rel_no, child_table, child_col, parent_table, parent_col]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = BORDER
            cell.alignment = WRAP
        row += 1

    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "64748B"
    for i, w in enumerate([5, 16, 24, 16, 24], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def create_workbook() -> Workbook:
    erd = load_erd(ERD_FILE)
    rows = parse_tables(erd)

    wb = Workbook()
    wb.remove(wb.active)

    build_index_sheet(wb.create_sheet("목차", 0), rows)
    build_definition_sheet(wb.create_sheet("테이블정의서", 1), rows)
    build_fk_sheet(wb.create_sheet("FK관계", 2), erd)

    return wb


def main() -> None:
    if not ERD_FILE.exists():
        raise FileNotFoundError(f"ERD file not found: {ERD_FILE}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = create_workbook()
    wb.save(OUT_FILE)

    erd = load_erd(ERD_FILE)
    table_count = len(erd["doc"]["tableIds"])
    col_count = len(erd["collections"]["tableColumnEntities"])
    fk_count = len(erd["collections"].get("relationshipEntities", {}))
    print(f"Generated: {OUT_FILE}")
    print(f"  Sheets: {', '.join(wb.sheetnames)}")
    print(f"  Tables: {table_count}, Columns: {col_count}, FK: {fk_count}")


if __name__ == "__main__":
    main()
