# -*- coding: utf-8 -*-
"""Excel 셀·서식 기반 상세 화면 목업."""
from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MOCKUP_COLS = 14  # A~N
MOCKUP_COL_MIN = 5.5
MOCKUP_COL_MAX = 12.0

# 그리드 헤더별 컴팩트 열 너비 (목업 전용)
GRID_COLUMN_WIDTHS: dict[str, float] = {
    "□": 5.5, "No": 5.5, "…": 5.5,
    "생성일": 8, "상신일": 8, "업무": 7, "저장": 6, "주의": 6,
    "Ref.No": 11, "프로세스": 8, "액티비티": 8, "고객번호": 10,
    "고객명": 10, "통화": 6, "금액": 9, "TotalText": 8, "항목심사": 8,
    "S/W": 6, "Depth": 6, "심사자": 8, "상태": 7, "유형": 7,
    "우선순위": 8, "회차": 6, "주기": 6, "DetectionID": 9,
    "마스터ID": 9, "접수일자": 8, "처리상태": 8, "제외사유": 8,
    "생성구분": 8, "접수자": 8, "담당자": 8, "진행구분": 8,
    "조건유형": 8, "연산자": 7, "조건값": 8, "사용": 5, "비고": 7,
    "직원명": 8, "직원번호": 9, "권한": 7, "역할": 6,
    "화면ID": 7, "화면명": 9, "접근": 5,
}

FILL_PRIMARY = PatternFill("solid", fgColor="1C5BBA")
FILL_PRIMARY_DK = PatternFill("solid", fgColor="144A9C")
FILL_HEADER = PatternFill("solid", fgColor="EFF6FF")
FILL_SURFACE = PatternFill("solid", fgColor="FFFFFF")
FILL_MUTED = PatternFill("solid", fgColor="F1F5F9")
FILL_PAGE = PatternFill("solid", fgColor="E8EEF5")  # 화면 바깥(iframe 배경)
FILL_WARN = PatternFill("solid", fgColor="FEF3C7")
FILL_ALERT = PatternFill("solid", fgColor="FEE2E2")
FILL_CLEAN = PatternFill("solid", fgColor="D1FAE5")
FILL_BTN = PatternFill("solid", fgColor="E2E8F0")
FILL_BTN_PRI = PatternFill("solid", fgColor="1C5BBA")
FILL_BTN_SEC = PatternFill("solid", fgColor="FFFFFF")
FILL_GRID_HDR = PatternFill("solid", fgColor="1E3A5F")
FILL_GRID_ALT = PatternFill("solid", fgColor="F8FAFC")
FILL_INPUT = PatternFill("solid", fgColor="FFFFFF")
FILL_LINK = PatternFill("solid", fgColor="DBEAFE")
FILL_TAB_ON = PatternFill("solid", fgColor="1C5BBA")
FILL_TAB_OFF = PatternFill("solid", fgColor="CBD5E1")
FILL_TREE_SEL = PatternFill("solid", fgColor="DBEAFE")
FILL_CANVAS = PatternFill("solid", fgColor="F8FAFC")

FONT_TITLE = Font(bold=True, color="FFFFFF", size=11)
FONT_SUB = Font(bold=True, color="144A9C", size=9)
FONT_BOLD = Font(bold=True, size=9, color="1E293B")
FONT_NORM = Font(size=9, color="1E293B")
FONT_MUTED = Font(size=8, color="64748B")
FONT_LABEL = Font(size=8, color="475569", bold=True)
FONT_BTN = Font(size=9, color="1C5BBA", bold=True)
FONT_BTN_W = Font(size=9, color="FFFFFF", bold=True)
FONT_REQ = Font(size=8, color="EF4444", bold=True)
FONT_LINK = Font(size=9, color="1C5BBA", underline="single")
FONT_ALERT = Font(size=9, color="EF4444", bold=True)
FONT_CLEAN = Font(size=9, color="059669", bold=True)
FONT_GRID_HDR = Font(bold=True, size=8, color="FFFFFF")
FONT_INPUT = Font(size=9, color="94A3B8")

SIDE_THIN = Side(style="thin", color="CBD5E1")
SIDE_MED = Side(style="medium", color="94A3B8")
SIDE_INPUT = Side(style="thin", color="94A3B8")
BORDER = Border(left=SIDE_THIN, right=SIDE_THIN, top=SIDE_THIN, bottom=SIDE_THIN)
BORDER_OUTER = Border(left=SIDE_MED, right=SIDE_MED, top=SIDE_MED, bottom=SIDE_MED)
BORDER_INPUT = Border(left=SIDE_INPUT, right=SIDE_INPUT, top=SIDE_INPUT, bottom=SIDE_INPUT)
ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_R = Alignment(horizontal="right", vertical="center", wrap_text=True)
ALIGN_TL = Alignment(horizontal="left", vertical="top", wrap_text=True)

SAMPLE_BY_HEADER: dict[str, str] = {
    "No": "1", "□": "☑", "생성일": "07/12", "상신일": "07/11", "업무": "수출",
    "Ref.No": "FX-2026-001", "저장": "Y", "주의": "—", "프로세스": "심사",
    "액티비티": "1차심사", "고객번호": "1234567890", "고객명": "(주)삼성전자",
    "통화": "USD", "금액": "125,000", "TotalText": "Alert", "항목심사": "OK",
    "S/W": "Clean", "Depth": "2", "심사자": "홍길동", "상태": "대기",
    "유형": "심사", "우선순위": "P0", "회차": "1", "주기": "D",
    "DetectionID": "D001", "마스터ID": "M-10001", "□": "☑",
    "접수일자": "07/12", "처리상태": "미확인", "제외사유": "—",
    "생성구분": "정상", "접수자": "user01", "담당자": "홍길동",
    "진행구분": "진행", "우선순위": "1", "조건유형": "AMOUNT",
    "연산자": "GTE", "조건값": "100000", "사용": "Y", "비고": "—",
    "직원명": "홍길동", "직원번호": "E10001", "권한": "심사자",
    "역할": "S1", "화면ID": "1010", "화면명": "ToDoList", "접근": "Y",
}


def _style_range(ws: Worksheet, r1: int, c1: int, r2: int, c2: int, fill=None, font=None, align=None, border=None) -> None:
    b = border if border is not None else BORDER
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = b
            if fill:
                cell.fill = fill
            if font:
                cell.font = font
            if align:
                cell.alignment = align


def _block(ws: Worksheet, r1: int, c1: int, r2: int, c2: int, text: str,
           fill=None, font=None, align=None, border=None) -> None:
    if r2 > r1 or c2 > c1:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = ws.cell(row=r1, column=c1, value=text)
    if font:
        cell.font = font
    if align:
        cell.alignment = align
    _style_range(ws, r1, c1, r2, c2, fill=fill, font=font, align=align, border=border)


def _set_row_heights(ws: Worksheet, row: int, height: float = 18) -> None:
    ws.row_dimensions[row].height = height


def _paint_page(ws: Worksheet, r1: int, r2: int, c1: int = 1, c2: int = MOCKUP_COLS) -> None:
    """iframe/페이지 배경."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            if cell.fill is None or cell.fill.fgColor is None or cell.fill.fgColor.rgb in (None, "00000000"):
                cell.fill = FILL_PAGE


def _btn(ws: Worksheet, row: int, c1: int, c2: int, label: str, *, primary: bool = False) -> None:
    if primary:
        _block(ws, row, c1, row, c2, label, FILL_BTN_PRI, FONT_BTN_W, ALIGN_C, BORDER)
    else:
        _block(ws, row, c1, row, c2, label, FILL_BTN_SEC, FONT_BTN, ALIGN_C, BORDER_INPUT)


def _input_field(ws: Worksheet, row: int, c1: int, c2: int, placeholder: str = "입력…") -> None:
    _block(ws, row, c1, row, c2, placeholder, FILL_INPUT, FONT_INPUT, ALIGN_L, BORDER_INPUT)


def _set_col_widths(ws: Worksheet, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def display_width(value: Any) -> int:
    """한글/전각 문자를 고려한 Excel 열 너비 산정용 표시 폭."""
    if value is None:
        return 0
    text = str(value)
    lines = text.splitlines() or [text]
    width = 0
    for line in lines:
        line_width = sum(2 if ord(ch) > 127 else 1 for ch in line)
        width = max(width, line_width)
    return width


def _column_width_from_text(text: str, *, min_w: float = MOCKUP_COL_MIN, max_w: float = MOCKUP_COL_MAX) -> float:
    width = display_width(text)
    if not width:
        return min_w
    return min(max_w, max(min_w, width * 0.9 + 1))


def _set_mockup_col_width(ws: Worksheet, col: int, width: float) -> None:
    """목업 A~N 열 — 표 보정과 분리해 고정(덮어쓰기)."""
    width = min(MOCKUP_COL_MAX, max(MOCKUP_COL_MIN, width))
    ws.column_dimensions[get_column_letter(col)].width = width


def _apply_col_width(ws: Worksheet, col: int, width: float) -> None:
    """목업 내부 계산용 — 동일 열 내 최대값 누적."""
    letter = get_column_letter(col)
    current = ws.column_dimensions[letter].width or 0
    ws.column_dimensions[letter].width = max(current, min(MOCKUP_COL_MAX, width))


def _grid_col_width(header: str, sample: str) -> float:
    if header in GRID_COLUMN_WIDTHS:
        return GRID_COLUMN_WIDTHS[header]
    return min(
        MOCKUP_COL_MAX,
        max(MOCKUP_COL_MIN, _column_width_from_text(header) * 0.85, _column_width_from_text(sample) * 0.85),
    )


def _fit_grid_col_widths(ws: Worksheet, c1: int, headers: list[str], sid: str) -> None:
    grid_widths: dict[int, float] = getattr(ws, "_mockup_grid_widths", {})
    for i, header in enumerate(headers):
        sample = _sample_value(header if header != "…" else "...", 1)
        width = _grid_col_width(header, sample)
        col = c1 + i
        grid_widths[col] = max(grid_widths.get(col, 0), width)
    ws._mockup_grid_widths = grid_widths  # type: ignore[attr-defined]


def auto_fit_mockup_columns(
    ws: Worksheet,
    start_row: int,
    end_row: int,
    c1: int = 1,
    c2: int = MOCKUP_COLS,
    *,
    min_width: float = MOCKUP_COL_MIN,
    max_width: float = MOCKUP_COL_MAX,
) -> None:
    """목업 영역 A~N 열 너비를 컴팩트하게 고정한다 (표 데이터와 분리)."""
    if end_row < start_row:
        return

    col_widths: dict[int, float] = {}
    processed_merges: set[tuple[int, int, int, int]] = set()

    for mr in ws.merged_cells.ranges:
        if mr.max_row < start_row or mr.min_row > end_row:
            continue
        if mr.max_col < c1 or mr.min_col > c2:
            continue

        key = (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
        if key in processed_merges:
            continue
        processed_merges.add(key)

        anchor = ws.cell(row=mr.min_row, column=mr.min_col)
        if anchor.value is None:
            continue

        ncol = mr.max_col - mr.min_col + 1
        # 넓은 병합(헤더·페이징·검색행)은 열당 폭을 제한 — 전체 텍스트로 열이 커지지 않게
        if ncol >= 5:
            per_col = min(max_width, 9.0)
        elif ncol >= 3:
            per_col = min(max_width, _column_width_from_text(str(anchor.value), max_w=max_width * ncol) / ncol)
        else:
            per_col = _column_width_from_text(str(anchor.value), min_w=min_width, max_w=max_width)

        for col in range(max(mr.min_col, c1), min(mr.max_col, c2) + 1):
            col_widths[col] = max(col_widths.get(col, 0), per_col)

    for row in range(start_row, end_row + 1):
        for col in range(c1, c2 + 1):
            in_merge = any(
                mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col
                for mr in ws.merged_cells.ranges
                if mr.max_row >= start_row and mr.min_row <= end_row
            )
            if in_merge:
                continue
            value = ws.cell(row=row, column=col).value
            if value is None:
                continue
            col_widths[col] = max(
                col_widths.get(col, 0),
                _column_width_from_text(str(value), min_w=min_width, max_w=max_width),
            )

    grid_widths: dict[int, float] = getattr(ws, "_mockup_grid_widths", {})
    for col in range(c1, c2 + 1):
        width = grid_widths.get(col, col_widths.get(col, 8.0))
        _set_mockup_col_width(ws, col, width)


def _sample_value(header: str, row_idx: int) -> str:
    base = SAMPLE_BY_HEADER.get(header, "…")
    if header in ("No",) or header == "□":
        return str(row_idx)
    if header == "Ref.No" and row_idx > 1:
        return f"FX-2026-{row_idx:03d}"
    if header == "금액" and row_idx > 1:
        return f"{125000 * row_idx:,}"
    return base


def _cell_style_for_value(header: str, val: str, sid: str, row_idx: int, *, selected: bool = False):
    if selected and header not in ("Ref.No", "TotalText", "S/W", "항목심사"):
        return FILL_TREE_SEL, FONT_NORM
    if header == "Ref.No":
        return FILL_SURFACE if not selected else FILL_TREE_SEL, FONT_LINK
    if val in ("Alert", "30") or "Alert" in str(val):
        return FILL_ALERT, FONT_ALERT
    if val in ("Clean", "40", "OK") or "Clean" in str(val):
        return FILL_CLEAN, FONT_CLEAN
    if header == "우선순위" or (sid == "1010" and row_idx == 1):
        return FILL_WARN, FONT_BOLD
    if header == "□":
        return FILL_SURFACE if not selected else FILL_TREE_SEL, FONT_NORM
    return FILL_SURFACE, FONT_NORM


def _draw_search_panel(ws: Worksheet, row: int, search: list, c1: int, c2: int) -> int:
    if not search:
        return row
    btn_w = 2
    field_area = c2 - btn_w
    n = min(len(search), 6)
    per = max(2, (field_area - c1 + 1) // min(n, 4))
    label_row = row
    input_row = row + 1
    col = c1
    for s in search[:6]:
        label = s[0]
        param = s[1] if len(s) > 1 else ""
        ftype = s[2] if len(s) > 2 else "text"
        req = s[3] if len(s) > 3 else "N"
        if col > field_area - 2:
            break
        end = min(col + per - 1, field_area)
        req_m = " *" if req == "Y" else ""
        _block(ws, label_row, col, label_row, end, f"{label}{req_m}", FILL_MUTED, FONT_LABEL, ALIGN_L)
        if ftype == "select":
            _block(ws, input_row, col, input_row, end, "전체 ▾", FILL_INPUT, FONT_INPUT, ALIGN_L, BORDER_INPUT)
        elif ftype in ("date", "datetime"):
            _block(ws, input_row, col, input_row, end, "YYYY-MM-DD", FILL_INPUT, FONT_INPUT, ALIGN_L, BORDER_INPUT)
        else:
            _input_field(ws, input_row, col, end, "입력…")
        col = end + 1
    _block(ws, label_row, c2 - 1, input_row, c2 - 1, "초기화", FILL_BTN_SEC, FONT_BTN, ALIGN_C, BORDER_INPUT)
    _block(ws, label_row, c2, input_row, c2, "조회", FILL_BTN_PRI, FONT_BTN_W, ALIGN_C, BORDER)
    _set_row_heights(ws, label_row, 16)
    _set_row_heights(ws, input_row, 20)
    return input_row + 1


def _draw_toolbar(ws: Worksheet, row: int, buttons: list, c1: int, c2: int) -> int:
    btns = [str(b[1]) for b in (buttons or []) if len(b) >= 2 and "[" in str(b[1])]
    if not btns:
        return row
    span = max(1, (c2 - c1 + 1) // len(btns))
    col = c1
    for btn in btns[:8]:
        end = min(col + span - 1, c2)
        label = btn.strip("[]") if btn.startswith("[") else btn
        primary = "조회" in btn or "저장" in btn or "확인" in btn
        _btn(ws, row, col, end, label, primary=primary)
        col = end + 1
        if col > c2:
            break
    _set_row_heights(ws, row, 22)
    return row + 1


def _draw_grid(ws: Worksheet, row: int, columns: list, c1: int, c2: int, sid: str, data_rows: int = 3) -> int:
    headers = []
    for c in columns:
        h = c[0]
        if h.startswith("(") and "hidden" in h.lower():
            continue
        headers.append(h)
    if not headers:
        headers = ["□", "No", "생성일", "업무", "Ref.No", "상태"]

    max_cols = c2 - c1 + 1
    if len(headers) > max_cols:
        display_headers = headers[: max_cols - 1] + ["…"]
    else:
        display_headers = headers

    col = c1
    for h in display_headers:
        _block(ws, row, col, row, col, h, FILL_GRID_HDR, FONT_GRID_HDR, ALIGN_C)
        col += 1
    _fit_grid_col_widths(ws, c1, display_headers, sid)
    _set_row_heights(ws, row, 20)
    row += 1

    for ri in range(1, data_rows + 1):
        selected = ri == 1
        row_fill = FILL_TREE_SEL if selected else (FILL_GRID_ALT if ri % 2 == 0 else FILL_SURFACE)
        col = c1
        for h in display_headers:
            if h == "…":
                val = f"+{len(headers) - max_cols + 1}"
            elif h == "□":
                val = "☑" if selected else "□"
            else:
                val = _sample_value(h, ri)
            fill, font = _cell_style_for_value(h, val, sid, ri, selected=selected)
            if h not in ("Ref.No", "TotalText", "S/W", "항목심사", "□") and not (
                val in ("Alert", "30", "Clean", "40", "OK") or "Alert" in str(val) or "Clean" in str(val)
            ):
                fill = row_fill
            _block(ws, row, col, row, col, val, fill, font, ALIGN_C)
            col += 1
        _set_row_heights(ws, row, 18)
        row += 1

    _block(ws, row, c1, row, c2, "◀  1  2  3  …  10  ▶     총 128건  |  50건/페이지", FILL_MUTED, FONT_MUTED, ALIGN_C)
    _set_row_heights(ws, row, 18)
    return row + 2


def _header_bar(ws: Worksheet, row: int, sid: str, name: str, domain: str, url: str, status: str, c1: int, c2: int) -> int:
    status_tag = f"  [{status}]" if status else ""
    _block(ws, row, c1, row, c2, f"  {name}  [{sid}]{status_tag}  ", FILL_PRIMARY, FONT_TITLE, ALIGN_L)
    _set_row_heights(ws, row, 24)
    row += 1
    bc = f"{domain}  >  {name}  |  {url}"
    _block(ws, row, c1, row, c2, bc, FILL_MUTED, FONT_MUTED, ALIGN_L)
    _set_row_heights(ws, row, 16)
    return row + 1


def _draw_detail_tree(ws: Worksheet, row: int, c1: int, c2: int, items: list[str]) -> int:
    _block(ws, row, c1, row, c2, "문서분류", FILL_GRID_HDR, FONT_GRID_HDR, ALIGN_L)
    _set_row_heights(ws, row, 18)
    row += 1
    for i, item in enumerate(items):
        prefix = "▶ " if i == 0 else "  "
        fill = FILL_TREE_SEL if i == 0 else FILL_MUTED
        font = FONT_BOLD if i == 0 else FONT_NORM
        _block(ws, row, c1, row, c2, f"{prefix}{item}", fill, font, ALIGN_L)
        _set_row_heights(ws, row, 16)
        row += 1
    return row


def _draw_detail_result_grid(ws: Worksheet, row: int, c1: int, c2: int) -> int:
    headers = ["항목명", "AI결과", "신뢰도", "심사결과", "비고"]
    col_span = max(1, (c2 - c1 + 1) // len(headers))
    col = c1
    for h in headers:
        end = min(col + col_span - 1, c2)
        _block(ws, row, col, row, end, h, FILL_GRID_HDR, FONT_GRID_HDR, ALIGN_C)
        col = end + 1
    _set_row_heights(ws, row, 18)
    row += 1

    rows_data = [
        ("수출자", "40 Clean", "92%", "(입력)", ""),
        ("수입자", "30 Alert", "78%", "(입력)", ""),
        ("금액", "40", "88%", "(입력)", ""),
    ]
    for ri, cells in enumerate(rows_data, 1):
        col = c1
        for ci, (h, val) in enumerate(zip(headers, cells)):
            end = min(col + col_span - 1, c2)
            fill, font = _cell_style_for_value(h, val, "1011", ri)
            if ci == 3:
                _input_field(ws, row, col, end, val)
            else:
                _block(ws, row, col, row, end, val, fill, font, ALIGN_C if ci else ALIGN_L)
            col = end + 1
        _set_row_heights(ws, row, 18)
        row += 1

    _block(ws, row, c1, row, c2, "[+ 행추가]", FILL_LINK, FONT_LINK, ALIGN_L)
    _set_row_heights(ws, row, 16)
    return row + 1


def draw_list_mockup(ws: Worksheet, row: int, sid: str, name: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    domain, url, status = screen[2], screen[3], screen[5]
    row = _header_bar(ws, row, sid, name, domain, url, status, c1, c2)
    row = _draw_search_panel(ws, row, d.get("search") or [], c1, c2)
    row = _draw_toolbar(ws, row, d.get("buttons") or [], c1, c2)
    row = _draw_grid(ws, row, d.get("columns") or [], c1, c2, sid, data_rows=3)
    return row


def draw_detail_mockup(ws: Worksheet, row: int, sid: str, name: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    domain, url, status, jsp = screen[2], screen[3], screen[5], screen[6]
    row = _header_bar(ws, row, sid, name, domain, url, status, c1, c2)

    fields = d.get("fields") or []
    hdr_parts = []
    for f in fields[:4]:
        hdr_parts.append(f"{f[0]}: sample")
    hdr_info = "  |  ".join(hdr_parts) or "Ref.No: FX-2026-001  |  업무: 수출  |  Depth: 1"
    _block(ws, row, c1, row, c2 - 1, f"상세 팝업 — {hdr_info}", FILL_PRIMARY_DK, FONT_TITLE, ALIGN_L)
    _block(ws, row, c2, row, c2, "✕", FILL_PRIMARY_DK, FONT_TITLE, ALIGN_C)
    _set_row_heights(ws, row, 22)
    row += 1

    l_w = 2
    m_w = 5
    l_end = c1 + l_w - 1
    m_end = l_end + m_w
    r_start = m_end + 1
    body_end = row + 5

    tree_items = ["선적서류", "P/L", "C/I", "B/L"]
    tree_end = _draw_detail_tree(ws, row, c1, l_end, tree_items)

    # 중앙 이미지 뷰어 (단일 블록 — 중첩 merge 금지)
    _block(
        ws, row, l_end + 1, body_end, m_end,
        "이미지 뷰어\n\n[문서 이미지]\n하이라이트 · 핀",
        FILL_CANVAS, FONT_MUTED, ALIGN_C, BORDER_OUTER,
    )

    tabs = ["항목심사", "TotalText/SW", "전달정보"]
    tab_w = max(1, (c2 - r_start + 1) // len(tabs))
    tc = r_start
    for i, tab in enumerate(tabs):
        te = tc + tab_w - 1 if i < len(tabs) - 1 else c2
        fill = FILL_TAB_ON if i == 0 else FILL_TAB_OFF
        font = FONT_BTN_W if i == 0 else FONT_BTN
        _block(ws, row, tc, row, te, tab, fill, font, ALIGN_C)
        tc = te + 1

    result_end = _draw_detail_result_grid(ws, row + 1, r_start, c2)
    row = max(tree_end, result_end, body_end + 1)

    btns = d.get("buttons") or []
    btn_labels = []
    for b in btns[:7]:
        label = str(b[1]).strip("[]") if str(b[1]).startswith("[") else str(b[1])
        btn_labels.append(label)
    if not btn_labels:
        btn_labels = ["임시저장", "결재상신", "최종승인", "재스캔", "B/L수정"]

    span = max(1, (c2 - c1 + 1) // len(btn_labels))
    col = c1
    for i, label in enumerate(btn_labels):
        end = min(col + span - 1, c2)
        primary = i in (1, 2) or "상신" in label or "승인" in label
        _btn(ws, row, col, end, label, primary=primary)
        col = end + 1
        if col > c2:
            break
    _set_row_heights(ws, row, 24)
    row += 1
    _block(ws, row, c1, row, c2, f"JSP: WEB-INF/jsp/{jsp}", FILL_MUTED, FONT_MUTED, ALIGN_L)
    return row + 2


def _draw_master_list(ws: Worksheet, row: int, c1: int, c2: int, users: list[tuple[str, str, str]]) -> int:
    _block(ws, row, c1, row, c2, "사용자 목록", FILL_GRID_HDR, FONT_GRID_HDR, ALIGN_L)
    _set_row_heights(ws, row, 18)
    row += 1
    _input_field(ws, row, c1, c2, "직원명·번호 검색…")
    _set_row_heights(ws, row, 18)
    row += 1
    for i, (name, emp_no, role) in enumerate(users):
        prefix = "▶ " if i == 0 else "  "
        fill = FILL_TREE_SEL if i == 0 else (FILL_GRID_ALT if i % 2 else FILL_SURFACE)
        _block(ws, row, c1, row, c2, f"{prefix}{name}  {emp_no}  {role}", fill, FONT_NORM, ALIGN_L)
        _set_row_heights(ws, row, 16)
        row += 1
    return row


def draw_master_detail_mockup(ws: Worksheet, row: int, sid: str, name: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    domain, url, status = screen[2], screen[3], screen[5]
    row = _header_bar(ws, row, sid, name, domain, url, status, c1, c2)
    row = _draw_search_panel(
        ws, row,
        d.get("search") or [("권한", "authCd", "select", "N", ""), ("사용여부", "usgYn", "select", "N", "")],
        c1, c2,
    )

    split = c1 + 4
    users = [
        ("홍길동", "E10001", "심사자"),
        ("이결재", "E10002", "결재자"),
        ("박QA", "E10003", "QA"),
        ("최관리", "E10004", "관리자"),
    ]
    master_end = _draw_master_list(ws, row, c1, split - 1, users)

    tabs = ["업무별 역할", "화면별 접근"]
    tc = split
    for i, tab in enumerate(tabs):
        te = split + 3 if i == 0 else c2
        fill = FILL_TAB_ON if i == 0 else FILL_TAB_OFF
        font = FONT_BTN_W if i == 0 else FONT_BTN
        _block(ws, row, tc, row, te, tab, fill, font, ALIGN_C)
        if i == 0:
            tc = te + 1

    fields = d.get("fields") or []
    if fields:
        detail_headers = [f[0] for f in fields[:6]]
    else:
        detail_headers = ["업무", "역할", "유효시작", "유효종료", "사용"]

    col = split
    col_span = max(1, (c2 - split + 1) // len(detail_headers))
    for h in detail_headers:
        end = min(col + col_span - 1, c2)
        _block(ws, row + 1, col, row + 1, end, h, FILL_GRID_HDR, FONT_GRID_HDR, ALIGN_C)
        col = end + 1
    _set_row_heights(ws, row + 1, 18)

    sample_vals = []
    for f in fields[:6] if fields else []:
        lbl = f[0]
        if "업무" in lbl:
            sample_vals.append("수출")
        elif "역할" in lbl:
            sample_vals.append("S1")
        elif "화면" in lbl:
            sample_vals.append("1010")
        elif "접근" in lbl:
            sample_vals.append("Y")
        elif "시작" in lbl or "종료" in lbl:
            sample_vals.append("2026-01-01" if "시작" in lbl else "9999-12-31")
        else:
            sample_vals.append("…")
    if not sample_vals:
        sample_vals = ["수출", "S1", "2026-01-01", "9999-12-31", "Y"]

    col = split
    for val in sample_vals:
        end = min(col + col_span - 1, c2)
        _block(ws, row + 2, col, row + 2, end, val, FILL_SURFACE, FONT_NORM, ALIGN_C)
        col = end + 1
    _set_row_heights(ws, row + 2, 18)

    _block(ws, row + 3, split, row + 3, c2, "[+ 추가]   ※ S1·S2 동시불가", FILL_LINK, FONT_LINK, ALIGN_L)
    detail_end = row + 4

    cols_detail = d.get("columns")
    if cols_detail:
        detail_end = _draw_grid(ws, row + 4, cols_detail, split, c2, sid, data_rows=2)

    row = max(master_end, detail_end)
    _block(ws, row, split, row, c2 - 2, "이력보기 9080", FILL_LINK, FONT_LINK, ALIGN_L)
    _btn(ws, row, c2 - 1, c2 - 1, "취소")
    _btn(ws, row, c2, c2, "저장", primary=True)
    return row + 2


def draw_dashboard_mockup(ws: Worksheet, row: int, sid: str, name: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    row = _header_bar(ws, row, sid, name, screen[2], screen[3], screen[5], c1, c2)

    _block(ws, row, c1, row, c2 - 2, "긴급 심사대상  |  RENEGO  |  TotalText Alert 30  |  SLA 30분", FILL_WARN, FONT_BOLD, ALIGN_L)
    _btn(ws, row, c2 - 1, c2, "심사하기", primary=True)
    _set_row_heights(ws, row, 28)
    row += 1

    _btn(ws, row, c2 - 1, c2, "새로고침")
    _set_row_heights(ws, row, 20)
    row += 1

    q = max(2, (c2 - c1 + 1) // 4)
    cards = [("심사 ToDo", "12", FILL_LINK), ("결재 ToDo", "3", FILL_SURFACE), ("QA ToDo", "5", FILL_SURFACE), ("긴급", "2", FILL_ALERT)]
    col = c1
    for title, cnt, fill in cards:
        end = min(col + q - 1, c2)
        _block(ws, row, col, row + 2, end, f"{title}\n\n{cnt}건", fill, FONT_BOLD, ALIGN_C, BORDER_OUTER)
        col = end + 1
    row += 3

    mid = (c1 + c2) // 2
    _block(ws, row, c1, row + 3, mid, "최근 처리 현황\n(간트/타임라인)", FILL_CANVAS, FONT_NORM, ALIGN_C, BORDER_OUTER)
    _block(ws, row, mid + 1, row + 3, c2, "AI 검출 추이\n(스파크라인)", FILL_CANVAS, FONT_NORM, ALIGN_C, BORDER_OUTER)
    row += 4
    _block(ws, row, c1, row, c2, "※ v2: v1 지도·도넛 차트 제거 (FR-001)", FILL_MUTED, FONT_MUTED, ALIGN_L)
    return row + 2


def draw_shell_mockup(ws: Worksheet, row: int, sid: str, name: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    # GNB: 구간별 셀 (전체 병합 후 겹침 쓰기 금지)
    mid = c1 + 5
    _block(ws, row, c1, row, mid, " 우리은행  |  AI Sanction", FILL_PRIMARY, FONT_TITLE, ALIGN_L)
    _block(ws, row, mid + 1, row, mid + 3, "DASHBOARD", FILL_PRIMARY_DK, FONT_BTN_W, ALIGN_C)
    _block(ws, row, mid + 4, row, c2 - 2, "심사 | 결재 | 통계", FILL_PRIMARY_DK, FONT_BTN_W, ALIGN_C)
    _block(ws, row, c2 - 1, row, c2, "홍길동", FILL_PRIMARY_DK, FONT_BTN_W, ALIGN_C)
    _set_row_heights(ws, row, 24)
    row += 1

    lnb = c1 + 3
    lnb_items = [
        ("심사", True),
        ("  ▶ ToDo  12", False),
        ("  진행상태", False),
        ("결재", True),
        ("  ToDo  3", False),
        ("위젯", True),
        ("  AI현황  428", False),
        ("  배지 ON", False),
    ]
    for i, (text, bold) in enumerate(lnb_items):
        fill = FILL_TREE_SEL if "▶" in text else FILL_MUTED
        font = FONT_BOLD if bold else FONT_NORM
        _block(ws, row + i, c1, row + i, lnb - 1, text, fill, font, ALIGN_L)
        _set_row_heights(ws, row + i, 16)

    _block(ws, row, lnb, row, lnb + 2, "[대시보드]", FILL_TAB_ON, FONT_BTN_W, ALIGN_C)
    _block(ws, row, lnb + 3, row, lnb + 6, "1010 ToDo ✕", FILL_TAB_OFF, FONT_BTN, ALIGN_C)
    _block(ws, row, lnb + 7, row, c2, "2010 ✕", FILL_TAB_OFF, FONT_BTN, ALIGN_C)
    _set_row_heights(ws, row, 20)

    iframe_end = row + 7
    _block(
        ws, row + 1, lnb, iframe_end, c2,
        "\n iframe (wf-view-container)\n\n 활성 화면 JSP 로드\n SSE 단일 연결 — layout.jsp\n",
        FILL_CANVAS, FONT_MUTED, ALIGN_C, BORDER_OUTER,
    )
    return iframe_end + 2


def draw_login_mockup(ws: Worksheet, row: int, sid: str, name: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    start = row
    pad = 3
    card_c1 = c1 + pad
    card_c2 = c2 - pad
    card_row = row + 1

    _paint_page(ws, start, start + 14, c1, c2)

    _block(ws, card_row, card_c1, card_row, card_c2, "우리은행 AI Sanction", FILL_SURFACE, FONT_BOLD, ALIGN_C, BORDER_OUTER)
    _set_row_heights(ws, card_row, 22)
    card_row += 1
    _block(ws, card_row, card_c1, card_row, card_c2, "AI 심사자동화 시스템", FILL_SURFACE, FONT_SUB, ALIGN_C, BORDER)
    _set_row_heights(ws, card_row, 18)
    card_row += 1

    fields = d.get("fields") or [("Username", "text", "Y", ""), ("Password", "password", "Y", "")]
    for f in fields:
        label = f[0]
        ftype = f[1] if len(f) > 1 else "text"
        req = " *" if len(f) > 2 and f[2] == "Y" else ""
        _block(ws, card_row, card_c1, card_row, card_c2, f"{label}{req}", FILL_SURFACE, FONT_LABEL, ALIGN_L, BORDER)
        _set_row_heights(ws, card_row, 16)
        card_row += 1
        ph = "●●●●●●" if ftype == "password" else "입력…"
        _input_field(ws, card_row, card_c1, card_c2, ph)
        _set_row_heights(ws, card_row, 20)
        card_row += 1

    _btn(ws, card_row, card_c1, card_c1 + 2, "로그인", primary=True)
    _btn(ws, card_row, card_c1 + 3, card_c2, "SSO 로그인")
    _set_row_heights(ws, card_row, 24)
    card_row += 1

    _style_range(ws, card_row - len(fields) * 2 - 3, card_c1, card_row, card_c2, fill=FILL_SURFACE, border=BORDER_OUTER)
    return card_row + 2


def draw_modal_mockup(ws: Worksheet, row: int, sid: str, name: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    row = _header_bar(ws, row, sid, name, screen[2], screen[3], screen[5], c1, c2)
    inner1, inner2 = c1 + 2, c2 - 2
    modal_start = row

    _block(ws, row, inner1, row, inner2 - 1, f"  {name}  [{sid}]", FILL_PRIMARY, FONT_TITLE, ALIGN_L)
    _block(ws, row, inner2, row, inner2, "✕", FILL_PRIMARY, FONT_TITLE, ALIGN_C)
    _set_row_heights(ws, row, 22)
    row += 1

    fields = d.get("fields") or []
    if fields:
        for f in fields[:6]:
            label = f[0]
            req = " *" if len(f) > 2 and f[2] == "Y" else ""
            ftype = f[1] if len(f) > 1 else "text"
            _block(ws, row, inner1, row, inner1 + 2, f"{label}{req}", FILL_MUTED, FONT_LABEL, ALIGN_L)
            if ftype == "select":
                _block(ws, row, inner1 + 3, row, inner2, "선택 ▾", FILL_INPUT, FONT_INPUT, ALIGN_L, BORDER_INPUT)
            else:
                _input_field(ws, row, inner1 + 3, inner2, "입력…" if len(f) <= 2 or f[2] != "Y" else "필수")
            _set_row_heights(ws, row, 20)
            row += 1
    else:
        _block(ws, row, inner1, row, inner2, "선택 건  N건", FILL_MUTED, FONT_BOLD, ALIGN_L)
        _set_row_heights(ws, row, 18)
        row += 1
        _block(ws, row, inner1, row, inner2, "일괄 승인 의견", FILL_MUTED, FONT_LABEL, ALIGN_L)
        _set_row_heights(ws, row, 16)
        row += 1
        _input_field(ws, row, inner1, inner2, "의견 입력…")
        _set_row_heights(ws, row, 36)
        row += 1
        _block(ws, row, inner1, row, inner2, "※ 확인 후 처리", FILL_MUTED, FONT_MUTED, ALIGN_L)
        _set_row_heights(ws, row, 16)
        row += 1

    _btn(ws, row, inner2 - 2, inner2 - 1, "취소")
    _btn(ws, row, inner2, inner2, "확인", primary=True)
    _set_row_heights(ws, row, 22)
    row += 1

    _style_range(ws, modal_start, inner1, row - 1, inner2, fill=FILL_SURFACE, border=BORDER_OUTER)
    _paint_page(ws, modal_start - 1, row, c1, c2)
    return row + 1


def draw_form_mockup(ws: Worksheet, row: int, sid: str, name: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    row = _header_bar(ws, row, sid, name, screen[2], screen[3], screen[5], c1, c2)
    fields = d.get("fields") or [("설정항목", "text", "Y", "")]
    for f in fields[:8]:
        label = f[0]
        ftype = f[1] if len(f) > 1 else "text"
        req = " *" if len(f) > 2 and f[2] == "Y" else ""
        db = f[3] if len(f) > 3 else ""
        _block(ws, row, c1, row, c1 + 2, f"{label}{req}", FILL_MUTED, FONT_LABEL, ALIGN_L)
        if ftype == "select":
            _block(ws, row, c1 + 3, row, c2 - (2 if db else 0), "선택 ▾", FILL_INPUT, FONT_INPUT, ALIGN_L, BORDER_INPUT)
        elif "date" in ftype:
            _block(ws, row, c1 + 3, row, c2 - (2 if db else 0), "YYYY-MM-DD", FILL_INPUT, FONT_INPUT, ALIGN_L, BORDER_INPUT)
        else:
            _input_field(ws, row, c1 + 3, c2 - (2 if db else 0), "입력값")
        if db:
            _block(ws, row, c2 - 1, row, c2, db[:12], FILL_MUTED, FONT_MUTED, ALIGN_C)
        _set_row_heights(ws, row, 22)
        row += 1
    row = _draw_toolbar(ws, row, d.get("buttons") or [("—", "[저장]", "POST")], c1, c2)
    return row


def draw_chart_mockup(ws: Worksheet, row: int, sid: str, name: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    row = _header_bar(ws, row, sid, name, screen[2], screen[3], screen[5], c1, c2)
    row = _draw_search_panel(ws, row, d.get("search") or [], c1, c2)
    mid = (c1 + c2) // 2
    _block(
        ws, row, c1, row + 5, mid,
        "Chart.js\n(막대/선/도넛)\n\n통계 차트 영역",
        FILL_CANVAS, FONT_NORM, ALIGN_C, BORDER_OUTER,
    )
    if d.get("columns"):
        row = _draw_grid(ws, row, d["columns"], mid + 1, c2, sid, data_rows=2)
    else:
        _block(
            ws, row, mid + 1, row + 5, c2,
            "집계 테이블\n(요약)",
            FILL_MUTED, FONT_NORM, ALIGN_C, BORDER_OUTER,
        )
        row += 6
    return row


def draw_generic_mockup(ws: Worksheet, row: int, sid: str, name: str, pattern: str, d: dict, screen: tuple) -> int:
    c1, c2 = 1, MOCKUP_COLS
    row = _header_bar(ws, row, sid, name, screen[2], screen[3], screen[5], c1, c2)
    for comp in (d.get("ui_components") or [])[:6]:
        _block(ws, row, c1, row, c1 + 2, str(comp[0]), FILL_MUTED, FONT_LABEL, ALIGN_L)
        _block(ws, row, c1 + 3, row, c2 - 2, str(comp[1]), FILL_SURFACE, FONT_NORM, ALIGN_L)
        _block(ws, row, c2 - 1, row, c2, str(comp[2])[:8], FILL_MUTED, FONT_MUTED, ALIGN_C)
        _set_row_heights(ws, row, 20)
        row += 1
    if d.get("search"):
        row = _draw_search_panel(ws, row, d["search"], c1, c2)
    if d.get("columns"):
        row = _draw_grid(ws, row, d["columns"], c1, c2, sid, data_rows=2)
    return row + 1


DRAWERS: dict[str, Any] = {
    "C": draw_list_mockup,
    "차트": draw_chart_mockup,
    "D": draw_detail_mockup,
    "E": draw_master_detail_mockup,
    "B": draw_dashboard_mockup,
    "A": draw_shell_mockup,
    "단독": draw_login_mockup,
    "G": draw_modal_mockup,
    "F": draw_form_mockup,
    "H": draw_list_mockup,
    "전역": draw_modal_mockup,
}


def draw_excel_cell_mockup(
    ws: Worksheet,
    start_row: int,
    sid: str,
    name: str,
    pattern: str,
    d: dict,
    screen: tuple | None = None,
) -> int:
    if screen is None:
        screen = (sid, name, "", "", pattern, "", "", "", "", "", "", "")
    drawer = DRAWERS.get(pattern)
    if drawer is None:
        end_row = draw_generic_mockup(ws, start_row, sid, name, pattern, d, screen)
    else:
        end_row = drawer(ws, start_row, sid, name, d, screen)
    return end_row
