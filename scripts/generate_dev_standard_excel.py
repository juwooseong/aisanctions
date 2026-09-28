# -*- coding: utf-8 -*-
"""개발 표준 가이드 Excel 생성 (상세판)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from dev_standard_data import (
    ASIS_ROLES,
    ASIS_TOBE_AUTH,
    BACKEND_LAYERS,
    BTNFLAG_MAP,
    COMMON_CODE_GROUPS,
    COMMON_LAYOUTS,
    DEPTH_LOGIC,
    FR_TRACE,
    FRONTEND_PATTERNS,
    IMPLEMENTATION_STEPS,
    NAMING_CONVENTIONS,
    PR_CHECKLIST,
    SESSION_FIELDS,
    SHEET_SECTIONS,
    SOD_RULES,
    TABLE_EXTENSIONS,
    TASK_DEFINITIONS,
    TOBE_CORE_SCREENS,
    TOBE_TABLES,
)
from generate_screen_design_excel import SCREENS
from screen_design_dev_data import BACKEND_MAP, PATTERN_IMPL
from screen_design_enrich import EVENTS_SPEC_NOTE, PATTERN_EVENTS
from screen_design_tables import TABLE_CATALOG

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "screen-design" / "excel"
OUT_FILE = OUT_DIR / "개발표준가이드.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1C5BBA")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
SECTION_FILL = PatternFill("solid", fgColor="EFF6FF")
SECTION_FONT = Font(bold=True, size=11, color="144A9C")
TITLE_FONT = Font(bold=True, size=14, color="1C5BBA")
META_FONT = Font(size=10, color="475569")
NOTE_FONT = Font(italic=True, size=9, color="64748B")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

SHEET_INDEX = [
    ("01_개발자가이드", "시트 16섹션·구현 순서·네이밍·PR 체크리스트"),
    ("02_화면ID체계", "ID 대역·패턴·상태·btnFlag"),
    ("03_UI표준", "디자인 토큰·CSS·레이아웃·버튼·접근성"),
    ("04_패턴별구현", "패턴 A~H 구현·이벤트·체크리스트"),
    ("05_권한API", "AS-IS/TO-BE·L1/L2/L3·SoD·API"),
    ("06_메시지검증", "MSG 코드·Confirm·화면별 검증"),
    ("07_전역UI", "SSE·Toast·LNB·postMessage"),
    ("08_참고문서", "와이어프레임·FR 추적·연계 문서"),
    ("09_화면카탈로그", "전체 62화면 목록 (도메인·패턴·상태)"),
    ("10_백엔드표준", "계층 구조·네이밍·핵심 화면 클래스"),
    ("11_프론트표준", "DataTables·공통 JS·JSP 패턴"),
    ("12_데이터모델", "TO-BE 테이블·확장 컬럼·공통코드"),
    ("13_TOBE핵심화면", "7050/7052/8070/6066~69 상세"),
]

PATTERN_DESC = {
    "C": "목록 — 검색+DataTables",
    "D": "상세팝업 — 3단분할",
    "E": "Master-Detail CRUD",
    "A": "메인 셸 — GNB+LNB+탭+iframe",
    "B": "대시보드 — 카드·긴급리스트",
    "G": "모달 — 소형 오버레이",
    "H": "이력 타임라인",
    "F": "폼/설정",
    "단독": "로그인 단독",
    "전역": "전역 컴포넌트",
    "차트": "차트 통계",
}


def write_section_title(ws, row: int, title: str, col_span: int = 6) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    cell = ws.cell(row=row, column=1, value=title)
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    cell.alignment = Alignment(vertical="center")
    for c in range(1, col_span + 1):
        ws.cell(row=row, column=c).border = BORDER
    return row + 1


def write_data_table(ws, row: int, headers: list[str], rows: list[tuple]) -> int:
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
        c.alignment = CENTER
    row += 1
    for data in rows:
        for col, val in enumerate(data, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = BORDER
            c.alignment = WRAP
        row += 1
    return row + 1


def write_text_block(ws, row: int, text: str, col_span: int = 6, *, font: Font | None = None) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.alignment = WRAP
    if font:
        cell.font = font
    for c in range(1, col_span + 1):
        ws.cell(row=row, column=c).border = BORDER
    return row + 1


def write_sheet_title(ws, title: str, col_span: int = 6) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_span)
    ws["A1"].value = title
    ws["A1"].font = TITLE_FONT
    return 3


def auto_fit_columns(ws, *, min_width: float = 8, max_width: float = 60) -> None:
    merged = {
        cell.coordinate
        for merged_range in ws.merged_cells.ranges
        for row in ws[merged_range.coord]
        for cell in row
    }
    max_by_col: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.coordinate in merged or cell.value is None:
                continue
            width = len(str(cell.value))
            max_by_col[cell.column] = max(max_by_col.get(cell.column, 0), width)
    for col_idx, width in max_by_col.items():
        letter = get_column_letter(col_idx)
        adjusted = min(max_width, max(min_width, width * 1.15 + 2))
        ws.column_dimensions[letter].width = adjusted


def build_cover_meta(ws) -> None:
    ws.merge_cells("A1:E1")
    ws["A1"].value = "개발 표준 가이드 (TO-BE) — 상세판"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A2:E2")
    ws["A2"].value = f"우리은행 AI 심사자동화 시스템 UI · 작성일 {date.today().isoformat()}"
    ws["A2"].font = META_FONT
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A3:E3")
    ws["A3"].value = (
        "출처: docs/screen-design/00_*.md · screen-design-specification.md · "
        "task-centric-auth-approval-guide.md · screen_design_dev_data.py"
    )
    ws["A3"].font = NOTE_FONT
    ws["A3"].alignment = Alignment(horizontal="center", wrap_text=True)


def build_index_sheet(ws) -> None:
    ws.title = "목차"
    build_cover_meta(ws)
    row = 5
    row = write_section_title(ws, row, "시트 목록", 5)
    index_rows = []
    for i, (name, desc) in enumerate(SHEET_INDEX, 1):
        if "01_" in name:
            target, src = "전체 개발자", "00_개발자_가이드.md"
        elif "02_" in name:
            target, src = "신규 담당자", "screen-design-specification.md §1.4~1.6"
        elif "03_" in name:
            target, src = "UI/프론트", "00_UI표준.md · spec §2~3"
        elif "04_" in name:
            target, src = "구현 담당", "screen_design_dev_data.py · screen_design_enrich.py"
        elif "05_" in name:
            target, src = "권한·백엔드", "00_권한_API_매트릭스.md · task-centric-auth-approval-guide.md"
        elif "06_" in name:
            target, src = "검증·QA", "00_메시지_검증.md · spec §9"
        elif "07_" in name:
            target, src = "셸·알림", "screen-design-specification.md §7"
        elif "08_" in name:
            target, src = "전체", "screen-design-specification.md §10"
        elif "09_" in name:
            target, src = "화면 담당", "generate_screen_design_excel.py SCREENS"
        elif "10_" in name:
            target, src = "백엔드", "screen_design_dev_data.py BACKEND_MAP"
        elif "11_" in name:
            target, src = "프론트", "screen_design_enrich.py PATTERN_EVENTS"
        elif "12_" in name:
            target, src = "DBA·백엔드", "task-centric-auth-approval-guide.md §4~5"
        else:
            target, src = "TO-BE 신규", "screen-design-specification.md §6"
        index_rows.append((i, name, desc, target, src))
    write_data_table(ws, row, ["No", "시트명", "설명", "대상", "원본 문서"], index_rows)
    row = write_section_title(ws, row, "문서 범위", 5)
    write_data_table(ws, row, ["포함", "제외"], [
        ("약 50개 비즈니스 화면 + TO-BE 신규", "TA 엔진·Flask AI 내부 UI"),
        ("v2 HYUNA 앱 셸 (GNB/LNB/탭)", "모바일 전용 화면"),
        ("와이어프레임 HTML 목업 연계", "인쇄용 보고서 레이아웃"),
    ])
    row = write_section_title(ws, row, "재생성 명령", 5)
    write_text_block(ws, row, "python scripts/generate_dev_standard_excel.py", 5)
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 36
    ws.freeze_panes = "A6"
    ws.sheet_properties.tabColor = "144A9C"


def build_dev_guide_sheet(ws) -> None:
    ws.title = "01_개발자가이드"
    row = write_sheet_title(ws, "개발자 가이드 — 화면설계서 시트 활용법")

    row = write_section_title(ws, row, "1. 문서 목적")
    row = write_text_block(
        ws, row,
        "화면별 시트(sheets/*.md)를 구현 명세로 사용. "
        "화면 구조·필드·이벤트·API·권한을 개발·검수·UAT 시 공통 기준으로 적용.",
    )

    row = write_section_title(ws, row, "2. 시트 구조 (16개 섹션) — 상세")
    row = write_data_table(ws, row, ["§", "섹션", "개발 활용", "구현 포인트", "산출물"], list(SHEET_SECTIONS))

    row = write_section_title(ws, row, "3. 구현 순서 (권장) — 단계별 산출물")
    row = write_data_table(ws, row, ["순서", "단계", "작업 내용", "구현 포인트", "산출물"], list(IMPLEMENTATION_STEPS))

    row = write_section_title(ws, row, "4. 패턴별 Quick Reference")
    row = write_data_table(ws, row, ["패턴", "대표 화면", "핵심 파일", "API 패턴", "권한"], [
        ("C 목록", "1010, 2010, 3010", "list.jsp + datatable.jsp + *ApiController", "GET /api/.../todo", "Task-Role + screenAccess"),
        ("D 상세", "1011, 2011, 3011", "detail.jsp + CommonDetailApiController", "POST /api/common/detail/*", "본인배정·SoD"),
        ("E MD", "7050, 7030, 8020", "master-detail JS + CRUD API", "GET/POST CRUD", "mgpeYn 또는 Role"),
        ("A 셸", "MAIN", "layout.jsp + shell-v2.js + notify-client.js", "SSE /api/notify/stream", "screenAccess"),
        ("B 대시", "DASH-v2", "dashboard.jsp + POST /api/dashboard", "POST", "Task-Role 숏컷"),
        ("G 모달", "BUNDLE, 1011-BL", "팝업 JSP + POST API", "POST", "부모 화면 권한"),
        ("H 이력", "9080, 9090", "history.jsp + GET API", "GET", "읽기전용"),
    ])

    row = write_section_title(ws, row, "5. 네이밍 규칙")
    row = write_data_table(ws, row, ["구분", "규칙", "예시", "경로"], list(NAMING_CONVENTIONS))

    row = write_section_title(ws, row, "6. PR·UAT 공통 체크리스트")
    for item in PR_CHECKLIST:
        row = write_text_block(ws, row, f"☐ {item}")

    row = write_section_title(ws, row, "7. 신규 화면 (코드 없음 — 우선 참조)")
    row = write_data_table(ws, row, ["ID", "URL", "FR", "우선 참조", "패턴"], [
        ("7050", "/admin/taskAuth", "FR-042, 054", "task-auth-panel.js, task-centric-auth-approval-guide.md", "E"),
        ("7052", "/admin/approvDepth", "FR-041, 043", "admin-7052-approv-depth.html", "E"),
        ("7051", "/admin/branch", "FR-056", "sheets/7051.md", "C"),
        ("8070", "/admin/aiMonitor", "FR-058, 064", "admin-8070-monitor.html", "C"),
        ("6066", "/stat/regulatory", "FR-066", "stats-6066-regulatory.html", "C"),
        ("6067", "/stat/efficiency", "FR-067", "stats-6067-efficiency.html", "C"),
        ("6069", "/stat/early-warning", "FR-069", "stats-6069-early-warning.html", "C"),
    ])

    row = write_section_title(ws, row, "8. 연계 문서·재생성")
    row = write_data_table(ws, row, ["리소스", "경로", "용도", "재생성"], [
        ("Markdown 시트", "docs/screen-design/sheets/*.md", "화면별 16섹션", "generate_screen_design_sheets.py"),
        ("Excel 통합본", "docs/screen-design/excel/화면설계서_TO-BE.xlsx", "68탭 화면설계", "generate_screen_design_excel.py"),
        ("프로세스 Excel", "docs/process/excel/업무프로세스정의서_TO-BE.xlsx", "76탭 프로세스", "generate_process_definition_excel.py"),
        ("본 가이드", "docs/screen-design/excel/개발표준가이드.xlsx", "개발 표준", "generate_dev_standard_excel.py"),
    ])
    ws.sheet_properties.tabColor = "1C5BBA"


def build_screen_id_sheet(ws) -> None:
    ws.title = "02_화면ID체계"
    row = write_sheet_title(ws, "화면 ID 체계·유형·상태", 5)

    row = write_section_title(ws, row, "1. 화면 ID 대역", 5)
    row = write_data_table(ws, row, ["대역", "도메인", "URL Prefix", "예시", "비고"], [
        ("LOGIN, MAIN, DASH", "공통 진입", "/login, /index, /dashboard", "LOGIN, MAIN, DASH-v2", "셸·대시보드"),
        ("1xxx", "심사", "/revert/*", "1010, 1011", "S1 Task-Role"),
        ("2xxx", "결재", "/app/*", "2010, 2011", "S2, Depth=2"),
        ("3xxx", "QA", "/qa/*", "3010, 3011", "QA Role"),
        ("4xxx", "심사 현황/명세", "/revert/stat/*", "4010~4050", "통계·명세"),
        ("5xxx", "업무일지", "/task/log/*", "5010~5040", "폼·목록"),
        ("6xxx", "통계", "/stat/*", "6010~6069", "차트·신규 통계"),
        ("7xxx", "관리자", "/admin/*", "7010, 7050, 7052", "사용자·권한·Rule"),
        ("8xxx", "관리자시스템", "/admin/*", "8010, 8070", "코드·시스템"),
        ("9xxx", "공통", "/common/*", "9080, 9090", "이력·팝업"),
    ])

    row = write_section_title(ws, row, "2. 화면 유형 (Pattern)", 5)
    row = write_data_table(ws, row, ["코드", "유형", "설명", "대표 화면", "CSS 핵심"], [
        ("A", "메인 셸", "GNB + LNB + 탭 + iframe", "MAIN", "wf-shell, wf-gnb, wf-lnb"),
        ("B", "대시보드", "카드·긴급 리스트", "DASH-v2", "wf-urgent-card, wf-shortcut-card"),
        ("C", "목록", "검색 + DataTables", "1010, 2010", "wf-search-v2, wf-table-v2"),
        ("D", "상세 팝업", "3단 분할 fullLayout", "1011, 2011", "wf-doc-tree, wf-image-viewer"),
        ("E", "Master-Detail", "좌우 2~3분할 CRUD", "7030, 7050", "wf-master-detail, wf-layer-tab"),
        ("F", "폼/설정", "단일 폼 + 저장", "5020, 7060", "wf-form"),
        ("G", "모달", "소형 오버레이", "BUNDLE, 1011-BL", "layer popup"),
        ("H", "이력", "타임라인 + 첨부", "9080", "wf-timeline"),
        ("단독", "로그인", "셸 없음", "LOGIN", "wf-login-card"),
        ("전역", "전역 컴포넌트", "layout.jsp 내", "GLOBAL", "wf-toast-urgent"),
        ("차트", "차트 통계", "검색 + Chart.js", "6010, 6030", "wf-chart-area"),
    ])

    row = write_section_title(ws, row, "3. 상태 범례", 5)
    row = write_data_table(ws, row, ["표기", "의미", "개발 시 주의", "UI", "권한"], [
        ("AS-IS", "현행 운영", "기존 코드 유지·점진 개선", "기존 CSS", "AI_INPT_AUT_CD"),
        ("v2", "HYUNA 셸·UI 개편", "wf-* CSS·shell-v2.js", "wireframe.css", "Task-Role 병행"),
        ("신규", "TO-BE 신규", "전체 신규 생성", "와이어프레임 1:1", "TO-BE 전용"),
        ("변경", "로직·UI 변경", "Task-Role/Depth/SoD", "v2 + 로직", "TO-BE 필터 추가"),
    ])

    row = write_section_title(ws, row, "4. btnFlag (상세 화면 역할 분기)", 5)
    row = write_data_table(ws, row, ["btnFlag", "역할", "화면ID", "버튼 세트", "API"], list(BTNFLAG_MAP))

    row = write_section_title(ws, row, "5. 화면 ID 주의사항", 5)
    row = write_text_block(
        ws, row,
        "※ 기존 7050(제재 Rule)과 TO-BE 7050(업무별 권한) 번호 충돌. "
        "운영 반영 시 기존 제재 Rule 화면 번호 재배정(예: 7053)을 DBA·업무와 확정.",
        5,
    )
    ws.sheet_properties.tabColor = "1C5BBA"


def build_ui_standard_sheet(ws) -> None:
    ws.title = "03_UI표준"
    row = write_sheet_title(ws, "UI/UX 설계 표준 (개발·목업 적용)", 5)

    row = write_section_title(ws, row, "1. 디자인 토큰", 5)
    row = write_data_table(ws, row, ["토큰", "값", "용도", "CSS 클래스", "적용 화면"], [
        ("Primary", "#1c5bba", "GNB, 버튼, 강조", "wf-gnb-active, .btn-primary", "전역"),
        ("Primary Dark", "#144a9c", "헤더 그라데이션", "—", "GNB"),
        ("Background", "#f8fafc", "페이지 배경", "body", "전역"),
        ("Surface", "#ffffff", "카드·패널", "wf-card", "대시·폼"),
        ("Border", "#cbd5e1", "구분선", "wf-border", "테이블·검색"),
        ("Text", "#1e293b", "본문", "—", "전역"),
        ("Muted", "#64748b", "보조 텍스트", "wf-muted", "breadcrumb"),
        ("Danger", "#ef4444", "Alert(30), 경보", "result-30", "목록·상세"),
        ("Success", "#10b981", "Clean(40), 정상", "result-40", "목록·상세"),
        ("Warning", "#d69e2e", "우선순위 강조 행", "wf-row-highlight", "1010 FR-004"),
    ])
    row = write_text_block(ws, row, "폰트: Noto Sans KR, Malgun Gothic — 13px 기본 | CSS: docs/wireframes/css/wireframe.css", 5, font=NOTE_FONT)

    row = write_section_title(ws, row, "2. 공통 CSS 클래스 (v2)", 5)
    row = write_data_table(ws, row, ["클래스", "용도", "패턴", "파일", "예시 화면"], [
        ("wf-search-v2", "검색 영역 (접기/펼치기)", "C", "wireframe.css", "1010, 2010"),
        ("wf-table-v2", "DataTables v2 스타일", "C", "wireframe.css", "전체 목록"),
        ("wf-paging", "페이징", "C", "wireframe.css", "전체 목록"),
        ("wf-row-highlight", "우선순위 행 (FR-004)", "C", "wireframe.css", "1010 P0/P1"),
        ("wf-master-detail", "Master-Detail 2분할", "E", "wireframe.css", "7050, 7030"),
        ("wf-layer-tab", "레이어 탭", "E", "wireframe.css", "7050"),
        ("wf-doc-tree", "문서 분류 트리", "D", "wireframe.css", "1011"),
        ("wf-image-viewer", "Canvas 이미지 뷰어", "D", "wireframe.css", "1011"),
        ("wf-result-tabs", "항목심사/TotalText 탭", "D", "wireframe.css", "1011"),
        ("wf-urgent-card", "긴급 심사 카드", "B", "wireframe.css", "DASH-v2"),
        ("wf-shortcut-card", "숏컷 카드", "B", "wireframe.css", "DASH-v2"),
        ("wf-badge-urgent", "URGENT 뱃지", "B/C", "wireframe.css", "1010, DASH"),
        ("wf-gnb-active", "GNB 활성 탭", "A", "wireframe.css", "MAIN"),
        ("wf-view-container", "iframe 컨테이너", "A", "wireframe.css", "MAIN"),
        ("wf-toast-urgent", "긴급 Toast", "전역", "wireframe.css", "GLOBAL"),
    ])

    row = write_section_title(ws, row, "3. 공통 레이아웃 컴포넌트", 5)
    row = write_data_table(ws, row, ["컴포넌트", "패턴", "구성", "파일", "TO-BE"], list(COMMON_LAYOUTS))

    row = write_section_title(ws, row, "4. MAIN 셸 영역 상세", 5)
    row = write_data_table(ws, row, ["영역", "구성", "TO-BE 변경", "JS", "이벤트"], [
        ("GNB", "DASHBOARD, 선적서류심사, WORKFLOW심사", "FR-002 2분류", "shell-v2.js", "탭 그룹 전환"),
        ("LNB", "메뉴 트리 + ToDo 배지", "FR-003 Task-Role 필터", "shell-v2.js", "ASIDE_REFRESH"),
        ("탭", "iframe MDI, 다중 탭", "기존 유지", "shell-v2.js", "openTab/closeTab"),
        ("SSE", "layout.jsp EventSource 1개", "FR-005 신규", "notify-client.js", "postMessage 분배"),
        ("Toast", "부모 body 전역", "notification-toast.js", "notification-toast.js", "URGENT"),
    ])

    row = write_section_title(ws, row, "5. Pattern C 목록 / Pattern D 상세 영역", 5)
    row = write_data_table(ws, row, ["패턴", "영역", "구성요소", "CSS", "비고"], [
        ("C", "헤더", "화면명 [ID] + breadcrumb", "—", ""),
        ("C", "검색", "조건 필드 + [초기화][조회]", "wf-search-v2", "우측 정렬"),
        ("C", "그리드", "DataTables 체크박스·정렬·페이징", "wf-table-v2", "serverSide"),
        ("C", "하단", "도메인별 액션 버튼", "—", ""),
        ("D", "헤더", "Ref.No, 업무, Depth 뱃지, [×]", "wf-badge", "TO-BE"),
        ("D", "B/L 탭", "복수 B/L 전환", "wf-bl-tab", ""),
        ("D", "좌측", "문서 분류 트리", "wf-doc-tree", ""),
        ("D", "중앙", "Canvas 이미지 뷰어", "wf-image-viewer", "핀·하이라이트"),
        ("D", "우측", "[항목심사][TotalText/SW][전달] 탭", "wf-result-tabs", ""),
        ("D", "하단", "역할·Depth별 가변 버튼", "—", "btnFlag·depth"),
    ])

    row = write_section_title(ws, row, "6. 버튼 규칙", 5)
    row = write_data_table(ws, row, ["유형", "스타일", "위치", "예시", "CSS"], [
        ("Primary", "파란 배경", "저장, 조회, 승인", "[조회] [저장] [최종승인]", ".btn-primary"),
        ("Secondary", "흰 배경 + 테두리", "취소, 초기화", "[초기화] [취소]", ".btn-secondary"),
        ("Danger", "빨간 테두리/배경", "삭제, Block", "[삭제]", ".btn-danger"),
        ("Icon", "fa-refresh 등", "검색 영역 우측", ".init_btn", "Font Awesome"),
    ])

    row = write_section_title(ws, row, "7. 결과 코드 색상", 5)
    row = write_data_table(ws, row, ["코드", "의미", "색상", "CSS/JS", "적용"], [
        ("30", "경보 (Alert)", "#ef4444", "result-30, 빨강", "TotalText·항목심사"),
        ("40", "정상 (Clean)", "#10b981", "result-40, 녹색", "TotalText·항목심사"),
        ("80", "수기입력 대상", "분홍 배경", "rowHandwriting", "inptAtvtCd=='80'"),
    ])

    row = write_section_title(ws, row, "8. 목업 유형", 5)
    row = write_data_table(ws, row, ["유형", "경로", "화면 수", "열기 방법", "비고"], [
        ("Standalone HTML", "docs/wireframes/*.html", "15+", "브라우저 직접 오픈", "1010v2, 1011v4"),
        ("허브 인라인", "index.html#screen-{ID}", "45+", "좌측 메뉴 또는 앵커", "index.html"),
        ("Excel 셀 목업", "화면설계서 §3.1", "62", "화면설계서_TO-BE.xlsx", "편집 가능"),
        ("패턴 템플릿", "UI 컴포넌트 표", "—", "본 가이드 §3", "목업 미작성"),
    ])

    row = write_section_title(ws, row, "9. 반응·접근성", 5)
    row = write_data_table(ws, row, ["항목", "규칙", "구현", "파일", ""], [
        ("해상도", "1920×1080 기준, 최소 1280px", "반응형 미지원", "wireframe.css", ""),
        ("브라우저", "Chrome, IE (레거시)", "SSE polyfill 검토", "notify-client.js", ""),
        ("키보드", "Tab: 검색 → 테이블 → 버튼", "tabindex", "JSP", ""),
        ("로딩", "Pace 인디케이터", "layout.jsp", "layout.jsp", ""),
        ("세션 만료", "ajaxError → 로그인", "fnLoginAlert()", "common.js", ""),
    ])
    ws.sheet_properties.tabColor = "1C5BBA"


def build_pattern_sheet(ws) -> None:
    ws.title = "04_패턴별구현"
    row = write_sheet_title(ws, "패턴별 구현 가이드 + 이벤트 명세", 6)
    row = write_text_block(ws, row, EVENTS_SPEC_NOTE, 6, font=NOTE_FONT)

    for code in ["C", "D", "E", "A", "B", "G", "H", "F", "단독", "전역", "차트"]:
        impl = PATTERN_IMPL.get(code, {})
        row = write_section_title(ws, row, f"패턴 {code} — {PATTERN_DESC.get(code, code)}", 6)

        screens = [f"{s[0]}({s[1]})" for s in SCREENS if s[4] == code][:6]
        if screens:
            row = write_text_block(ws, row, f"대표 화면: {', '.join(screens)}", 6, font=NOTE_FONT)

        layout = impl.get("layout", "")
        if layout:
            row = write_text_block(ws, row, layout.replace("```", "").strip(), 6)

        frontend = impl.get("frontend", [])
        if frontend:
            row = write_data_table(ws, row, ["구분", "파일/함수", "설명", "패턴", "비고", ""], [
                (r[0], r[1], r[2] if len(r) > 2 else "", code, "", "")
                for r in frontend
            ])

        backend = impl.get("backend", [])
        if backend:
            row = write_data_table(ws, row, ["계층", "구현", "비고", "패턴", "", ""], [
                (r[0], r[1], r[2] if len(r) > 2 else "", code, "", "")
                for r in backend
            ])

        events = PATTERN_EVENTS.get(code, [])
        if events:
            row = write_section_title(ws, row, f"이벤트 명세 ({code})", 6)
            row = write_data_table(ws, row, ["트리거", "동작", "비고", "구현", "", ""], [
                (e[0], e[1], e[2], "개발자 직접 구현", "", "")
                for e in events
            ])

        checklist = impl.get("checklist", [])
        if checklist:
            row = write_section_title(ws, row, f"체크리스트 ({code})", 6)
            for item in checklist:
                row = write_text_block(ws, row, f"☐ {item}", 6)
        row += 1

    ws.sheet_properties.tabColor = "1C5BBA"


def build_auth_api_sheet(ws) -> None:
    ws.title = "05_권한API"
    row = write_sheet_title(ws, "화면-API-권한 매트릭스 (TO-BE)", 6)

    row = write_section_title(ws, row, "1. AS-IS vs TO-BE 권한 모델", 6)
    row = write_data_table(ws, row, ["항목", "AS-IS", "TO-BE", "전환", "구현"], [
        (*r, "점진 전환" if i < 4 else "이중 읽기")
        for i, r in enumerate(ASIS_TOBE_AUTH)
    ])

    row = write_section_title(ws, row, "2. AS-IS 역할 코드", 6)
    row = write_data_table(ws, row, ["코드", "역할", "화면", "저장 위치", "TO-BE 대응"], [
        (r[0], r[1], r[2], r[3], "CSPD120TI S1/S2/QA")
        for r in ASIS_ROLES
    ])

    row = write_section_title(ws, row, "3. 업무(Task) 정의", 6)
    row = write_data_table(ws, row, ["TASK_CD", "BIZ_DSCD", "업무명", "업무군", "GRP"], list(TASK_DEFINITIONS))

    row = write_section_title(ws, row, "4. L1/L2/L3 검증 위치", 6)
    row = write_data_table(ws, row, ["계층", "테이블/엔진", "검증 시점", "구현 위치", "서비스", "비고"], [
        ("L1 Task-Role", "CSPD120TI", "ToDo 조회·상세 진입", "Service/Mapper WHERE", "TaskAuthService", "S1/S2/QA"),
        ("L1 Screen", "CSPD122TI", "메뉴·URL 진입", "AuthenticationInterceptor", "Interceptor", "screenAccess[]"),
        ("L2 Depth", "CSPD121TI", "건 생성·상세 로드", "DepthRuleEngine", "ApprovalDepthService", "1=단독, 2=교차"),
        ("L3 SoD", "비즈니스 규칙", "상신·승인", "Service", "AppDetailService 등", "S1≠S2"),
    ])

    row = write_section_title(ws, row, "5. SoD(직무 분리) 규칙", 6)
    row = write_data_table(ws, row, ["유형", "규칙", "검증 시점", "메시지/구현", "화면", ""], [
        (r[0], r[1], r[2], r[3], r[3].split()[-1] if "MSG" in r[3] else "", "")
        for r in SOD_RULES
    ])

    row = write_section_title(ws, row, "6. 세션 필드 (TO-BE)", 6)
    row = write_data_table(ws, row, ["필드", "설명", "용도", "DB 매핑", "", ""], [
        (r[0], r[1], r[2], r[3], "", "")
        for r in SESSION_FIELDS
    ])

    row = write_section_title(ws, row, "7. 승인 Depth 로직", 6)
    row = write_data_table(ws, row, ["단계", "내용", "시점", "구현", "화면", ""], [
        (r[0], r[1], r[2], r[3], "", "")
        for r in DEPTH_LOGIC
    ])

    row = write_section_title(ws, row, "8. 화면별 접근 (TO-BE)", 6)
    row = write_data_table(ws, row, ["화면", "L1 Task-Role", "L1 화면접근", "추가 조건", "API", "기존"], [
        ("1010", "S1 (해당 업무)", "1010 ∈ screenAccess", "AI_INSPE_ENO=loginUser", "GET /api/revert/todo", "01 OR"),
        ("1011", "S1 + 본인 배정", "1011", "AI_INSPE_ENO=loginUser", "POST /api/common/detail/*", "01"),
        ("2010", "S2 (해당 업무)", "2010", "AI_INPT_APPV_DEPTH_CD=2", "GET /api/app/todo", "02"),
        ("2011", "S2 + Depth=2", "2011", "S1≠S2 (SoD)", "POST /api/common/appDetail/*", "02"),
        ("3010", "QA (해당 업무)", "3010", "QA Role", "GET /api/qa/todo", "03"),
        ("7050", "—", "7050", "mgpeYn=Y", "GET/POST /api/admin/taskAuth", "관리자"),
        ("7052", "—", "7052", "mgpeYn=Y", "GET/POST /api/admin/approvDepth", "관리자"),
        ("8070", "—", "8070", "mgpeYn=Y", "GET /api/admin/aiMonitor", "관리자"),
    ])
    row = write_text_block(ws, row, "메뉴 표시: CSPD109TI AND screenAccess AND (해당 업무 Role 존재)", 6, font=NOTE_FONT)

    row = write_section_title(ws, row, "9. 핵심 API 매핑", 6)
    row = write_data_table(ws, row, ["화면", "API", "Method", "Request 주요", "Response", "비고"], [
        ("1010", "/api/revert/todo", "GET", "SearchVO, pageIndex", "resultList, paginationInfo", "S1 ToDo"),
        ("1011", "/api/common/detail/load", "POST", "inptMstSrno, inptAtvtCd", "마스터·항목·이미지", "상세 로드"),
        ("1011", "/api/common/revertDetail/santionSave", "POST", "DetailVO", "T200", "결재상신"),
        ("1011", "/api/common/revertDetail/santionApprv", "POST", "DetailVO", "T200", "최종승인"),
        ("2010", "/api/app/todo", "GET", "SearchVO", "resultList", "S2 ToDo"),
        ("2010", "/api/app/todo/bundle", "POST", "선택 건 목록", "T200", "일괄승인"),
        ("DASH", "/api/dashboard", "POST", "—", "긴급·숏컷", "대시보드"),
        ("7050", "/api/admin/taskAuth", "GET, POST", "TaskRoleVO", "목록/저장결과", "권한 설정"),
        ("7052", "/api/admin/approvDepth", "GET, POST", "ApprovDepthRuleVO", "규칙 목록", "Depth 규칙"),
        ("7052", "/api/admin/approvDepth/simulate", "POST", "업무·금액·Alert", "예상 Depth", "시뮬레이션"),
        ("LOGIN", "/login", "POST", "userId, password", "redirect /index", "세션 적재"),
        ("LOGIN", "/sso/prx", "POST", "SSO 토큰", "redirect /index", "FR-057"),
        ("전역", "/api/notify/stream", "GET (SSE)", "—", "이벤트 스트림", "SSE"),
        ("전역", "/api/notify/poll", "GET", "lastEventId", "이벤트 목록", "fallback"),
        ("내부", "/api/internal/notify/assign", "POST", "배정 정보", "—", "배정 알림"),
    ])
    ws.sheet_properties.tabColor = "1C5BBA"


def build_message_sheet(ws) -> None:
    ws.title = "06_메시지검증"
    row = write_sheet_title(ws, "유효성·메시지 (개발 구현)", 5)

    row = write_section_title(ws, row, "1. 공통 메시지 코드", 5)
    row = write_data_table(ws, row, ["코드", "메시지", "구현 위치", "화면", "처리"], [
        ("MSG-SESSION-001", "세션이 만료되었습니다. 다시 로그인하세요.", "CommonExceptionHandler", "전역", "로그인 리다이렉트"),
        ("MSG-AUTH-001", "해당 화면에 접근 권한이 없습니다.", "AuthenticationInterceptor", "전역", "접근 차단"),
        ("MSG-SOD-001", "동일 업무에 S1·S2 역할을 동시에 부여할 수 없습니다.", "TaskAuthService.save", "7050", "저장 차단"),
        ("MSG-SOD-002", "본인이 처리한 건은 결재할 수 없습니다.", "AppDetailService", "2011", "승인 차단"),
        ("MSG-DEPTH-001", "단독결재 건입니다. 최종승인으로 처리됩니다.", "1011 confirm", "1011", "confirm 후 승인"),
        ("MSG-NOTIFY-001", "긴급 심사가 배정되었습니다.", "notification-toast.js", "Toast", "알림 표시"),
    ])

    row = write_section_title(ws, row, "2. Confirm 다이얼로그", 5)
    row = write_data_table(ws, row, ["화면", "동작", "조건", "메시지", "API"], [
        ("1011", "최종승인", "depth=1", "단독결재로 최종 승인합니다. 계속하시겠습니까?", "santionApprv"),
        ("1011", "결재상신", "depth=2", "결재자에게 상신합니다. 계속하시겠습니까?", "santionSave"),
        ("2010", "일괄승인", "선택 N건", "선택 N건을 일괄 승인합니다.", "bundle"),
        ("7050", "역할 삭제", "행 선택", "해당 역할을 삭제하시겠습니까?", "taskAuth delete"),
        ("7052", "규칙 삭제", "선택 삭제", "선택 규칙을 삭제하시겠습니까?", "USG_YN=N"),
        ("전역", "저장", "폼 변경", "저장하시겠습니까?", "—"),
    ])

    row = write_section_title(ws, row, "3. 화면별 검증 규칙", 5)
    row = write_data_table(ws, row, ["화면", "검증", "시점", "구현", "메시지"], [
        ("7050", "S1+S2 동시 부여 금지", "저장", "Service UK 검증", "MSG-SOD-001"),
        ("7050", "필수: 업무, 역할, 유효기간", "저장", "beanValidator", "필수값 오류"),
        ("7050", "화면접근 중복", "저장", "CSPD122TI UK", "중복 경고"),
        ("7052", "우선순위 중복 경고", "저장", "Service", "경고 confirm"),
        ("7052", "DEFAULT 규칙 1건", "저장", "Service", "필수 규칙"),
        ("1011", "본인 배정 건만 처리", "진입", "Service WHERE", "MSG-AUTH-001"),
        ("1011", "depth=1 → 최종승인 확인", "클릭", "JS confirm", "MSG-DEPTH-001"),
        ("2010", "depth=2 건만 목록", "조회", "Mapper WHERE", "—"),
        ("2011", "S1≠S2 (SoD)", "승인", "Service", "MSG-SOD-002"),
        ("LOGIN", "Username/Password 필수", "제출", "beanValidator", "필수값"),
    ])

    row = write_section_title(ws, row, "4. 프론트 공통 검증", 5)
    row = write_data_table(ws, row, ["함수/패턴", "용도", "호출 시점", "파일", "비고"], [
        ("checkLogin()", "세션 유효성", "ajax 호출 전", "common.js", "false 시 중단"),
        ("fnLoginAlert()", "세션 만료 알림", "checkLogin 실패", "common.js", "로그인 이동"),
        ("beanValidator", "서버 폼 검증", "POST 요청", "Spring Validator", "서버 측"),
        ("ajaxError handler", "401/403 처리", "ajax 실패", "common.js", "세션 만료"),
        ("confirm()", "사용자 확인", "저장·삭제·승인", "각 화면 JS", "메시지 코드 연동"),
    ])
    row = write_text_block(ws, row, "if (!checkLogin()) { fnLoginAlert(); return false; }", 5)
    ws.sheet_properties.tabColor = "1C5BBA"


def build_global_ui_sheet(ws) -> None:
    ws.title = "07_전역UI"
    row = write_sheet_title(ws, "전역 UI (알림·Toast·SSE)", 5)

    row = write_section_title(ws, row, "1. 긴급 알람 Toast (FR-005)", 5)
    row = write_data_table(ws, row, ["항목", "값", "파일", "비고", ""], [
        ("컴포넌트", "notification-toast.js", "docs/wireframes/js/", "신규"),
        ("와이어프레임", "etc-fr005-urgent-notify.html", "docs/wireframes/", ""),
        ("렌더 위치", "layout.jsp 부모 document.body", "layout.jsp", "iframe 외부"),
        ("트리거", "SSE URGENT 이벤트", "notify-client.js", ""),
        ("[심사하기]", "1011 팝업 1-Step 진입", "detailUrl 파라미터", ""),
        ("[나중에]", "Toast 닫기, LNB 배지 유지", "—", ""),
    ])
    row = write_text_block(
        ws, row,
        "┌─────────────────────────────────────────────┐\n"
        "│ ⚠ 긴급 심사 배정                    [×]    │\n"
        "│ RENEGO · Ref.FX-2026-001                   │\n"
        "│ TotalText Alert 30 · SLA 30분              │\n"
        "│ [나중에]              [심사하기 →]         │\n"
        "└─────────────────────────────────────────────┘",
        5,
    )

    row = write_section_title(ws, row, "2. SSE 클라이언트", 5)
    row = write_data_table(ws, row, ["항목", "값", "파일", "비고", ""], [
        ("JS (신규)", "notify-client.js", "resources/js/", "layout.jsp include"),
        ("연결", "GET /api/notify/stream", "NotifyApiController", "EventSource 1회"),
        ("Fallback", "GET /api/notify/poll", "NotifyApiController", "SSE 미지원 시"),
        ("분배", "postMessage → 활성 iframe", "notify-client.js", "타입별 라우팅"),
        ("재연결", "onerror → exponential backoff", "notify-client.js", "최대 재시도"),
    ])

    row = write_section_title(ws, row, "3. SSE 이벤트 상세", 5)
    row = write_data_table(ws, row, ["이벤트", "수신", "동작", "페이로드", "FR"], [
        ("URGENT", "부모 Toast", "긴급 알람 표시", "refNo, taskNm, reason, detailUrl", "FR-005"),
        ("TODO_UPDATE", "1010/2010/3010", "DataTables reload", "taskCd (선택)", "FR-003"),
        ("DASHBOARD_REFRESH", "DASH", "POST /api/dashboard", "—", "FR-001"),
        ("ASIDE_REFRESH", "LNB", "배지 건수 갱신", "badgeCounts{}", "FR-003"),
    ])

    row = write_section_title(ws, row, "4. LNB 배지·설정 (FR-003)", 5)
    row = write_data_table(ws, row, ["요소", "설명", "데이터 소스", "갱신", "설정"], [
        ("업무별 건수", "수출 12, 수입 8 등", "ToDo API 집계", "ASIDE_REFRESH", "—"),
        ("AI 할당 ON/OFF", "자동 배정 수신", "사용자 설정", "—", "로컬 저장"),
        ("알림 ON/OFF", "긴급 알람 수신", "사용자 설정", "—", "FR-005"),
        ("숏컷 카드", "심사/결재 ToDo", "dashboard API", "DASHBOARD_REFRESH", "DASH-v2"),
    ])

    row = write_section_title(ws, row, "5. postMessage 프로토콜", 5)
    row = write_data_table(ws, row, ["type", "발신", "수신", "data", "동작"], [
        ("TODO_UPDATE", "notify-client (부모)", "iframe 1010/2010", "{taskCd}", "tableReload()"),
        ("DASHBOARD_REFRESH", "notify-client", "iframe DASH", "{}", "dashboard reload"),
        ("ASIDE_REFRESH", "notify-client", "shell-v2 LNB", "{badges}", "배지 갱신"),
        ("TAB_OPEN", "shell-v2", "—", "{url, title}", "탭 오픈"),
    ])
    ws.sheet_properties.tabColor = "1C5BBA"


def build_reference_sheet(ws) -> None:
    ws.title = "08_참고문서"
    row = write_sheet_title(ws, "참고 문서·FR 추적", 4)

    row = write_section_title(ws, row, "1. 와이어프레임·리소스", 4)
    row = write_data_table(ws, row, ["리소스", "경로", "용도"], [
        ("HTML 허브", "docs/wireframes/index.html", "50+ 화면 인라인 목업"),
        ("공통 CSS", "docs/wireframes/css/wireframe.css", "wf-* 클래스"),
        ("앱 셸 JS", "docs/wireframes/js/shell-v2.js", "탭·iframe"),
        ("7050 JS", "docs/wireframes/js/task-auth-panel.js", "업무별 권한"),
        ("Toast JS", "docs/wireframes/js/notification-toast.js", "긴급 알람"),
        ("목업 미리보기", "docs/screen-design/excel/previews/", "PNG (선택)"),
    ])

    row = write_section_title(ws, row, "2. 관련 문서", 4)
    row = write_data_table(ws, row, ["문서", "경로", "용도"], [
        ("화면설계서", "docs/screen-design-specification.md", "TO-BE 화면 설계 전체"),
        ("개발자 가이드", "docs/screen-design/00_개발자_가이드.md", "시트 읽는 법"),
        ("UI 표준", "docs/screen-design/00_UI표준.md", "디자인·패턴"),
        ("권한/API", "docs/screen-design/00_권한_API_매트릭스.md", "권한·API"),
        ("메시지", "docs/screen-design/00_메시지_검증.md", "메시지·검증"),
        ("권한/결재 가이드", "docs/task-centric-auth-approval-guide.md", "7050/7052/Depth/SoD"),
        ("아키텍처", "docs/tobe-architecture-design.md", "TO-BE 아키텍처"),
        ("와이어프레임", "docs/wireframes.md", "ASCII 와이어프레임"),
        ("프로세스", "docs/process/*.md", "화면별 프로세스·API"),
        ("BPMN", "bpmn/*.bpmn", "BPMN 2.0 (52개)"),
        ("FR 매핑", "docs/requirements-wireframe-map.md", "FR 추적"),
        ("테이블 정의", "docs/table-design/excel/테이블정의서.xlsx", "ERD 기반"),
    ])

    row = write_section_title(ws, row, "3. FR 추적 (상세)", 4)
    row = write_data_table(ws, row, ["FR", "화면/대상", "내용"], list(FR_TRACE))
    ws.sheet_properties.tabColor = "1C5BBA"


def build_catalog_sheet(ws) -> None:
    ws.title = "09_화면카탈로그"
    row = write_sheet_title(ws, f"전체 화면 카탈로그 ({len(SCREENS)}개)", 8)
    row = write_data_table(ws, row, ["No", "ID", "화면명", "도메인", "URL", "패턴", "상태", "와이어프레임"], [
        (i, s[0], s[1], s[2], s[3], s[4], s[5], s[8] or "—")
        for i, s in enumerate(SCREENS, 1)
    ])
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = "1C5BBA"


def build_backend_sheet(ws) -> None:
    ws.title = "10_백엔드표준"
    row = write_sheet_title(ws, "백엔드 개발 표준", 5)

    row = write_section_title(ws, row, "1. 계층 구조", 5)
    row = write_data_table(ws, row, ["계층", "구분", "어노테이션", "역할"], list(BACKEND_LAYERS))

    row = write_section_title(ws, row, "2. 네이밍 규칙", 5)
    row = write_data_table(ws, row, ["구분", "규칙", "예시", "경로"], list(NAMING_CONVENTIONS))

    row = write_section_title(ws, row, "3. API 응답 표준", 5)
    row = write_data_table(ws, row, ["항목", "규칙", "예시", "비고", ""], [
        ("성공 코드", "T200", '{"resultCode":"T200"}', "공통", ""),
        ("목록 응답", "resultList + paginationInfo", "HashMap", "DataTables", ""),
        ("오류", "CommonExceptionHandler", "MSG-* 코드", "전역", ""),
        ("SearchVO", "@ModelAttribute 바인딩", "검색 파라미터", "GET API", ""),
        ("페이징", "PaginationInfo", "pageIndex, recordCountPerPage", "serverSide", ""),
    ])

    row = write_section_title(ws, row, "4. MyBatis XML 규칙", 5)
    row = write_data_table(ws, row, ["항목", "규칙", "예시", "비고", ""], [
        ("파일명", "{Mapper}_SQL.xml", "RevertTodo_SQL.xml", "mappers/", ""),
        ("namespace", "Mapper Interface FQCN", "egovframework...RevertTodoMapper", "", ""),
        ("selectList", "목록 조회", "id=selectList", "권한 WHERE", ""),
        ("selectListTotCnt", "총 건수", "id=selectListTotCnt", "페이징", ""),
        ("동적 SQL", "<if test=...>", "검색 조건", "SearchVO", ""),
    ])

    row = write_section_title(ws, row, "5. 핵심 화면 백엔드 클래스", 5)
    rows = []
    for sid, info in BACKEND_MAP.items():
        rows.append((
            sid,
            info.get("controller", "—"),
            info.get("api", "—"),
            info.get("service", "—"),
            info.get("mapper", "—"),
        ))
    row = write_data_table(ws, row, ["화면ID", "Controller", "ApiController", "Service", "Mapper"], rows)
    ws.sheet_properties.tabColor = "1C5BBA"


def build_frontend_sheet(ws) -> None:
    ws.title = "11_프론트표준"
    row = write_sheet_title(ws, "프론트엔드 개발 표준", 5)

    row = write_section_title(ws, row, "1. 공통 JS 패턴", 5)
    row = write_data_table(ws, row, ["패턴", "함수/설정", "설명", "적용", "파일"], [
        (r[0], r[1], r[2], r[3], "common.js / JSP")
        for r in FRONTEND_PATTERNS
    ])

    row = write_section_title(ws, row, "2. DataTables 설정 템플릿", 5)
    row = write_text_block(
        ws, row,
        "var dataTable_XXXXConfig = {\n"
        "  ajaxUrl: '/api/...',\n"
        "  serverSide: true,\n"
        "  columns: [{ data: 'fieldName', ... }],\n"
        "  getSearchOption: function() { return $('#searchForm').serialize(); }\n"
        "};",
        5,
    )

    row = write_section_title(ws, row, "3. JSP 구조 (Pattern C)", 5)
    row = write_data_table(ws, row, ["영역", "파일/태그", "내용", "비고", ""], [
        ("레이아웃", "blank/ 뷰", "iframe 내 로드", "탭 MDI", ""),
        ("검색 폼", "<form id='searchForm'>", "wf-search-v2", "id/name = 시트 §4", ""),
        ("그리드", "datatable.jsp include", "DataTables div", "공통", ""),
        ("hidden form", "formDetail", "fnFullWin 대상", "상세 진입", ""),
        ("스크립트", "dataTable_XXXConfig", "columns, events", "JSP 하단", ""),
    ])

    row = write_section_title(ws, row, "4. JSP 구조 (Pattern D)", 5)
    row = write_data_table(ws, row, ["영역", "구현", "API", "비고", ""], [
        ("진입", "POST hidden form", "/revert/detail", "fnFullWin", ""),
        ("로드", "onLoad AJAX", "POST /api/common/detail/load", "마스터·항목", ""),
        ("이미지", "Canvas + annotation-store", "이미지 URL", "좌표 핀", ""),
        ("저장", "버튼 click", "inspection/save/apprv", "depth 분기", ""),
        ("닫기", "opener.tableReload()", "—", "부모 갱신", ""),
    ])

    row = write_section_title(ws, row, "5. CSS 적용 규칙", 5)
    row = write_data_table(ws, row, ["규칙", "내용", "예시", "비고", ""], [
        ("v2 화면", "wf-* 클래스 필수", "wf-search-v2", "AS-IS→v2 전환", ""),
        ("결과 색상", "30=빨강, 40=녹색", "fnRowCallback", "목록·상세", ""),
        ("우선순위", "wf-row-highlight", "P0/P1 행", "FR-004", ""),
        ("수기입력", "rowHandwriting", "inptAtvtCd==80", "1010", ""),
    ])
    ws.sheet_properties.tabColor = "1C5BBA"


def build_datamodel_sheet(ws) -> None:
    ws.title = "12_데이터모델"
    row = write_sheet_title(ws, "TO-BE 데이터 모델", 5)

    row = write_section_title(ws, row, "1. 신규 테이블 (TO-BE)", 5)
    row = write_data_table(ws, row, ["테이블", "논리명", "PK/UK", "FR", "비고"], list(TOBE_TABLES))

    row = write_section_title(ws, row, "2. 기존 테이블 확장", 5)
    row = write_data_table(ws, row, ["테이블", "컬럼", "설명", "시점", ""], [
        (r[0], r[1], r[2], r[3], "")
        for r in TABLE_EXTENSIONS
    ])

    row = write_section_title(ws, row, "3. 공통코드 그룹 (신규)", 5)
    row = write_data_table(ws, row, ["그룹", "그룹명", "코드 예시", "사용처", ""], [
        (r[0], r[1], r[2], r[3], "")
        for r in COMMON_CODE_GROUPS
    ])

    row = write_section_title(ws, row, "4. 핵심 테이블 카탈로그 (전체)", 5)
    catalog = sorted(TABLE_CATALOG.items(), key=lambda x: x[0])
    row = write_data_table(ws, row, ["물리명", "논리명", "구분", "출처", ""], [
        (phys, logical, "TO-BE" if phys >= "CSPD119TI" and phys <= "CSPD123TH" else "AS-IS", "sql/01_ddl.sql", "")
        for phys, logical in catalog
    ])
    ws.sheet_properties.tabColor = "1C5BBA"


def build_tobe_core_sheet(ws) -> None:
    ws.title = "13_TOBE핵심화면"
    row = write_sheet_title(ws, "TO-BE 신규·핵심 변경 화면 상세", 6)

    for scr in TOBE_CORE_SCREENS:
        row = write_section_title(ws, row, f"SCR-{scr['id']} — {scr['name']}", 6)
        row = write_data_table(ws, row, ["항목", "값", "비고", "", "", ""], [
            ("화면ID", scr["id"], "신규" if scr["id"] in ("7050", "7052", "8070", "6066", "6067", "6069") else "", "", "", ""),
            ("URL", scr["url"], "", "", "", ""),
            ("FR", scr["fr"], "", "", "", ""),
            ("구성", scr["tabs"], "", "", "", ""),
            ("API", scr["api"], "", "", "", ""),
            ("테이블", scr["tables"], "", "", "", ""),
            ("검증", scr["validation"], "", "", "", ""),
            ("참조", scr["ref"], "", "", "", ""),
        ])
        row += 1

    row = write_section_title(ws, row, "7052 규칙 목록 컬럼", 6)
    row = write_data_table(ws, row, ["컬럼", "필드", "설명", "비고", "", ""], [
        ("□", "—", "선택", "", "", ""),
        ("우선순위", "AI_INPT_RULE_PRIO", "낮을수록 우선", "", "", ""),
        ("업무", "AI_INPT_TASK_CD", "NULL=공통", "", "", ""),
        ("조건유형", "AI_INPT_COND_TYPE", "AMOUNT, ALERT_YN, RISK_LEVEL, DEFAULT", "", "", ""),
        ("연산자", "AI_INPT_COND_OPERATOR", "LT, GTE, EQ, IN", "", "", ""),
        ("조건값", "AI_INPT_COND_VALUE", "", "", "", ""),
        ("Depth", "AI_INPT_APPROV_DEPTH", "1(단독) / 2(교차)", "", "", ""),
        ("사용", "AI_INPT_RULE_USG_YN", "Y/N", "", "", ""),
    ])

    row = write_section_title(ws, row, "1011 Depth별 버튼 분기", 6)
    row = write_data_table(ws, row, ["Depth", "버튼", "API", "confirm", "", ""], [
        ("1 (단독)", "[최종승인]", "santionApprv", "MSG-DEPTH-001", "", ""),
        ("2 (교차)", "[결재상신]", "santionSave", "결재상신 confirm", "", ""),
        ("공통", "[임시저장]", "inspection", "—", "", ""),
    ])
    ws.sheet_properties.tabColor = "1C5BBA"


def create_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    build_index_sheet(wb.create_sheet("목차", 0))
    build_dev_guide_sheet(wb.create_sheet("01_개발자가이드"))
    build_screen_id_sheet(wb.create_sheet("02_화면ID체계"))
    build_ui_standard_sheet(wb.create_sheet("03_UI표준"))
    build_pattern_sheet(wb.create_sheet("04_패턴별구현"))
    build_auth_api_sheet(wb.create_sheet("05_권한API"))
    build_message_sheet(wb.create_sheet("06_메시지검증"))
    build_global_ui_sheet(wb.create_sheet("07_전역UI"))
    build_reference_sheet(wb.create_sheet("08_참고문서"))
    build_catalog_sheet(wb.create_sheet("09_화면카탈로그"))
    build_backend_sheet(wb.create_sheet("10_백엔드표준"))
    build_frontend_sheet(wb.create_sheet("11_프론트표준"))
    build_datamodel_sheet(wb.create_sheet("12_데이터모델"))
    build_tobe_core_sheet(wb.create_sheet("13_TOBE핵심화면"))

    for ws in wb.worksheets:
        auto_fit_columns(ws)

    return wb


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = create_workbook()
    wb.save(OUT_FILE)
    print(f"Saved: {OUT_FILE} ({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
