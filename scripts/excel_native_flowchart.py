# -*- coding: utf-8 -*-
"""Excel COM 후처리용 네이티브 플로우차트 배치 정보 생성."""
from __future__ import annotations

import json
from collections import defaultdict, deque
from pathlib import Path
from typing import Any

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FLOW_COLS = 16
ANCHOR_COL = 3  # C열
FLOWCHART_LEGEND = (
    "범례  타원: 시작/종료  |  둥근사각형: 업무  |  마름모: 분기  |  "
    "진한 회색 화살표: 흐름  |  연한 회색 화살표: 되돌아감"
)
LEGEND_FONT = Font(size=9, color="64748B")
LEGEND_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
MANIFEST: list[dict[str, Any]] = []


def reset_manifest() -> None:
    MANIFEST.clear()


def save_manifest(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps({"charts": MANIFEST}, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )


def _steps_to_flow(steps: list) -> tuple[list, list]:
    elements = [("Start", "시작", "시작", "startEvent")]
    flows: list[tuple] = []
    prev = "Start"
    for i, step in enumerate(steps[:10], 1):
        name = str(step[1] if len(step) > 1 else f"단계{i}")[:28]
        eid = f"S{i}"
        elements.append((eid, "업무", name, "task"))
        flows.append((f"F{i}", prev, eid, "", ""))
        prev = eid
    elements.append(("End", "종료", "종료", "endEvent"))
    flows.append(("F_end", prev, "End", "", ""))
    return elements, flows


def _node_size(node_type: str) -> tuple[int, int]:
    if node_type in ("시작", "종료"):
        return 82, 42
    if node_type == "분기":
        return 112, 66
    return 166, 52


def _build_layout(elements: list, flows: list) -> dict[str, Any]:
    elem_map = {e[0]: e for e in elements}
    outgoing: dict[str, list[str]] = defaultdict(list)
    incoming: dict[str, list[str]] = defaultdict(list)
    for flow in flows:
        src, tgt = flow[1], flow[2]
        outgoing[src].append(tgt)
        incoming[tgt].append(src)

    start_id = next((e[0] for e in elements if e[1] == "시작"), None)
    levels: dict[str, int] = {}
    if start_id:
        levels[start_id] = 0
        queue: deque[str] = deque([start_id])
        while queue:
            src = queue.popleft()
            for tgt in outgoing.get(src, []):
                if tgt not in levels:
                    levels[tgt] = levels[src] + 1
                    queue.append(tgt)

    next_level = max(levels.values(), default=-1) + 1
    for eid in elem_map:
        if eid not in levels:
            levels[eid] = next_level
            next_level += 1

    by_level: dict[int, list[str]] = defaultdict(list)
    for eid, level in levels.items():
        by_level[level].append(eid)

    canvas_width = 920
    top_margin = 52
    level_gap = 112
    nodes: list[dict[str, Any]] = []
    positions: dict[str, dict[str, Any]] = {}

    for level in sorted(by_level):
        ids = by_level[level]
        count = len(ids)
        spacing = min(230, canvas_width // max(count, 1))
        total = spacing * (count - 1)
        center = canvas_width // 2
        for idx, eid in enumerate(ids):
            _, node_type, name, _ = elem_map[eid]
            width, height = _node_size(node_type)
            cx = center - total // 2 + idx * spacing
            node = {
                "id": eid,
                "type": node_type,
                "name": name,
                "x": cx - width // 2,
                "y": top_margin + level * level_gap,
                "w": width,
                "h": height,
            }
            nodes.append(node)
            positions[eid] = node

    edges: list[dict[str, Any]] = []
    for flow in flows:
        fid, src, tgt, label = flow[0], flow[1], flow[2], flow[3] or ""
        if src not in positions or tgt not in positions:
            continue
        edges.append({
            "id": fid,
            "source": src,
            "target": tgt,
            "label": label,
            "loop": levels[tgt] <= levels[src],
        })

    max_bottom = max((n["y"] + n["h"] for n in nodes), default=220)
    return {
        "width": canvas_width,
        "height": max_bottom + 55,
        "nodes": nodes,
        "edges": edges,
    }


def render_flowchart(
    ws: Worksheet,
    start_row: int,
    bpmn_data: dict[str, Any],
    *,
    steps_fallback: list | None = None,
    title: str = "",
    screen_id: str = "",
    cache_dir: Path | None = None,
) -> tuple[int, tuple[int, int]]:
    """시트 공간을 확보하고 COM 후처리용 배치 정보를 기록한다."""
    elements = list(bpmn_data.get("elements") or [])
    flows = list(bpmn_data.get("flows") or [])
    if not elements and steps_fallback:
        elements, flows = _steps_to_flow(steps_fallback)

    if not elements:
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=10)
        ws.cell(row=start_row, column=1, value="흐름 정의 없음")
        return start_row + 2, (start_row, start_row + 1)

    layout = _build_layout(elements, flows)
    rows_needed = max(16, int(layout["height"] / 16) + 5)
    for row in range(start_row, start_row + rows_needed):
        ws.row_dimensions[row].height = 16

    legend_row = start_row + rows_needed - 1
    ws.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=10)
    legend_cell = ws.cell(row=legend_row, column=1, value=FLOWCHART_LEGEND)
    legend_cell.font = LEGEND_FONT
    legend_cell.alignment = LEGEND_ALIGN

    MANIFEST.append({
        "sheet": ws.title,
        "screen_id": screen_id,
        "title": title or "업무 흐름도",
        "anchor": f"{get_column_letter(ANCHOR_COL)}{start_row}",
        **layout,
    })

    end_row = start_row + rows_needed
    return end_row + 1, (start_row, end_row)


def fit_flowchart_columns(ws: Worksheet, start_row: int, end_row: int) -> None:
    for col in range(1, FLOW_COLS + 1):
        ws.column_dimensions[get_column_letter(col)].width = 10
