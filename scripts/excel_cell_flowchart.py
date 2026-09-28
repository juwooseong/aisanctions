# -*- coding: utf-8 -*-
"""Excel 셀·서식 기반 업무 흐름도 (BPMN 표 대신 시각 레이아웃)."""
from __future__ import annotations

from collections import defaultdict
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FLOW_COLS = 16

FILL_START = PatternFill("solid", fgColor="22C55E")
FILL_END = PatternFill("solid", fgColor="64748B")
FILL_TASK = PatternFill("solid", fgColor="EDE9FE")
FILL_TASK_HDR = PatternFill("solid", fgColor="7C3AED")
FILL_GATE = PatternFill("solid", fgColor="FEF08A")
FILL_ARROW = PatternFill("solid", fgColor="FAFAFA")
FILL_BRANCH = PatternFill("solid", fgColor="F0FDF4")
FILL_LOOP = PatternFill("solid", fgColor="FFF7ED")
FILL_LEGEND = PatternFill("solid", fgColor="F8FAFC")

FONT_START = Font(bold=True, color="FFFFFF", size=10)
FONT_END = Font(bold=True, color="FFFFFF", size=10)
FONT_TASK = Font(size=9, color="1E293B", bold=True)
FONT_TASK_SUB = Font(size=8, color="5B21B6")
FONT_GATE = Font(bold=True, size=9, color="854D0E")
FONT_ARROW = Font(size=12, color="7C3AED", bold=True)
FONT_LABEL = Font(size=8, color="64748B")
FONT_BRANCH = Font(size=8, color="166534")
FONT_LOOP = Font(size=8, color="C2410C", italic=True)
FONT_LEGEND = Font(size=8, color="64748B")

SIDE = Side(style="thin", color="94A3B8")
SIDE_THICK = Side(style="medium", color="7C3AED")
BORDER = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)
BORDER_TASK = Border(left=SIDE_THICK, right=SIDE_THICK, top=SIDE_THICK, bottom=SIDE_THICK)

ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=True)

NODE_H = 2
COL_START = 3
COL_WIDTH = 12


def _set_row(ws: Worksheet, row: int, height: float = 18) -> None:
    ws.row_dimensions[row].height = height


def _write_row_bar(
    ws: Worksheet, row: int, text: str, *,
    fill: PatternFill, font: Font,
    align: Alignment = ALIGN_C,
    height: float = 18,
) -> int:
    ws.merge_cells(start_row=row, start_column=COL_START, end_row=row, end_column=COL_START + COL_WIDTH - 1)
    cell = ws.cell(row=row, column=COL_START, value=text)
    cell.fill = fill
    cell.font = font
    cell.alignment = align
    cell.border = BORDER
    for c in range(COL_START, COL_START + COL_WIDTH):
        ws.cell(row=row, column=c).border = BORDER
        ws.cell(row=row, column=c).fill = fill
    _set_row(ws, row, height)
    return row + 1


def _write_task(ws: Worksheet, row: int, title: str, sub: str = "") -> int:
    r1 = row
    r2 = row + NODE_H - 1
    ws.merge_cells(start_row=r1, start_column=COL_START, end_row=r1, end_column=COL_START + COL_WIDTH - 1)
    top = ws.cell(row=r1, column=COL_START, value=title)
    top.fill = FILL_TASK_HDR
    top.font = FONT_TASK_SUB
    top.alignment = ALIGN_C
    top.border = BORDER_TASK

    ws.merge_cells(start_row=r1 + 1, start_column=COL_START, end_row=r2, end_column=COL_START + COL_WIDTH - 1)
    body = ws.cell(row=r1 + 1, column=COL_START, value=sub or " ")
    body.fill = FILL_TASK
    body.font = FONT_TASK
    body.alignment = ALIGN_C
    body.border = BORDER_TASK

    for r in range(r1, r2 + 1):
        for c in range(COL_START, COL_START + COL_WIDTH):
            ws.cell(row=r, column=c).border = BORDER_TASK
        _set_row(ws, r, 17)
    return r2 + 1


def _arrow(ws: Worksheet, row: int, label: str = "") -> int:
    text = f"▼  {label}" if label else "▼"
    return _write_row_bar(ws, row, text, fill=FILL_ARROW, font=FONT_ARROW, height=14)


def _branch_line(ws: Worksheet, row: int, prefix: str, label: str, target: str, is_loop: bool = False) -> int:
    fill = FILL_LOOP if is_loop else FILL_BRANCH
    font = FONT_LOOP if is_loop else FONT_BRANCH
    text = f"{prefix} {label} → {target}" if label else f"{prefix} → {target}"
    if is_loop:
        text = f"↺ {text}"
    return _write_row_bar(ws, row, text, fill=fill, font=font, align=ALIGN_L, height=16)


def _steps_to_flow(steps: list) -> tuple[list, list]:
    elements = [("Start", "시작", "시작", "startEvent")]
    flows: list[tuple] = []
    prev = "Start"
    for i, step in enumerate(steps[:10], 1):
        name = str(step[1] if len(step) > 1 else f"단계{i}")[:30]
        eid = f"S{i}"
        elements.append((eid, "업무", name, "task"))
        flows.append((f"F{i}", prev, eid, "", ""))
        prev = eid
    elements.append(("End", "종료", "종료", "endEvent"))
    flows.append((f"F_end", prev, "End", "", ""))
    return elements, flows


def _build_graph(bpmn_data: dict[str, Any]):
    elements = bpmn_data.get("elements") or []
    flows = bpmn_data.get("flows") or []
    elem_map = {e[0]: e for e in elements}
    outgoing: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for _, src, tgt, label, _ in flows:
        outgoing[src].append((tgt, label or ""))
    start_id = next((e[0] for e in elements if e[1] == "시작"), None)
    return elem_map, outgoing, start_id


def _node_label(elem_map: dict, nid: str) -> str:
    if nid not in elem_map:
        return nid
    return elem_map[nid][2]


def _render_node(ws: Worksheet, row: int, elem_map: dict, nid: str) -> int:
    if nid not in elem_map:
        return row
    _, etype, name, _ = elem_map[nid]
    if etype == "시작":
        return _write_row_bar(ws, row, f"●  {name}", fill=FILL_START, font=FONT_START)
    if etype == "종료":
        return _write_row_bar(ws, row, f"■  {name}", fill=FILL_END, font=FONT_END)
    if etype == "분기":
        return _write_row_bar(ws, row, f"◆  {name}", fill=FILL_GATE, font=FONT_GATE)
    return _write_task(ws, row, f"□  {name}", nid)


def render_flowchart(
    ws: Worksheet,
    start_row: int,
    bpmn_data: dict[str, Any],
    *,
    steps_fallback: list | None = None,
    title: str = "",
) -> tuple[int, tuple[int, int]]:
    elements = bpmn_data.get("elements") or []
    flows = bpmn_data.get("flows") or []
    if not elements and steps_fallback:
        elements, flows = _steps_to_flow(steps_fallback)

    data = {"elements": elements, "flows": flows}
    elem_map, outgoing, start_id = _build_graph(data)

    row = start_row
    flow_start = row

    if title:
        row = _write_row_bar(ws, row, title, fill=FILL_TASK_HDR, font=Font(bold=True, color="FFFFFF", size=10))

    row = _write_row_bar(
        ws, row,
        "● 시작  ■ 종료  ◆ 분기  □ 업무  ▼ 순차흐름  ↺ 되돌아감",
        fill=FILL_LEGEND, font=FONT_LEGEND, height=14,
    )
    row += 1

    if not elem_map or not start_id:
        row = _write_row_bar(ws, row, "흐름 정의 없음", fill=FILL_ARROW, font=FONT_LABEL)
        return row + 1, (flow_start, row)

    visited: set[str] = set()

    def walk(nid: str, row_ptr: int, depth: int = 0) -> int:
        if nid in visited or depth > 30:
            return row_ptr
        visited.add(nid)

        row_ptr = _render_node(ws, row_ptr, elem_map, nid)
        children = outgoing.get(nid, [])
        if not children:
            return row_ptr

        etype = elem_map.get(nid, ("", "", "", ""))[1]

        if etype == "분기" and len(children) > 1:
            row_ptr = _arrow(ws, row_ptr, "")
            prefixes = ["┌─", "├─"] * (len(children) - 1) + ["└─"]
            for i, (child_id, label) in enumerate(children):
                prefix = prefixes[i] if i < len(prefixes) else "├─"
                target_name = _node_label(elem_map, child_id)
                already = child_id in visited
                row_ptr = _branch_line(ws, row_ptr, prefix, label, target_name, is_loop=already)

                if already:
                    continue

                if elem_map.get(child_id, ("", "", "", ""))[1] == "종료":
                    row_ptr = _arrow(ws, row_ptr, "")
                    row_ptr = _render_node(ws, row_ptr, elem_map, child_id)
                    visited.add(child_id)
                    continue

                row_ptr = _render_node(ws, row_ptr, elem_map, child_id)
                visited.add(child_id)
                for tgt2, lbl2 in outgoing.get(child_id, []):
                    if tgt2 in visited or tgt2 == nid:
                        row_ptr = _branch_line(
                            ws, row_ptr, "  ", lbl2 or "재처리",
                            _node_label(elem_map, tgt2), is_loop=True,
                        )
                    elif elem_map.get(tgt2, ("", "", "", ""))[1] == "종료":
                        row_ptr = _arrow(ws, row_ptr, "")
                        row_ptr = _render_node(ws, row_ptr, elem_map, tgt2)
                        visited.add(tgt2)
            return row_ptr

        for child_id, label in children:
            row_ptr = _arrow(ws, row_ptr, label)
            if child_id in visited:
                row_ptr = _branch_line(ws, row_ptr, "↺", label, _node_label(elem_map, child_id), is_loop=True)
                continue
            if elem_map.get(child_id, ("", "", "", ""))[1] == "종료":
                row_ptr = _render_node(ws, row_ptr, elem_map, child_id)
                visited.add(child_id)
            else:
                row_ptr = walk(child_id, row_ptr, depth + 1)
        return row_ptr

    row = walk(start_id, row)
    flow_end = row
    return flow_end + 1, (flow_start, flow_end)


def fit_flowchart_columns(ws: Worksheet, start_row: int, end_row: int) -> None:
    for c in range(1, FLOW_COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = 8 if c < COL_START else 10
